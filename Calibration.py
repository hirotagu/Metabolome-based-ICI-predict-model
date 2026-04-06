"""
Calibration support analysis from saved prediction files
=======================================================
This script reads already-saved prediction CSVs produced by the existing
analysis scripts and computes supportive calibration metrics:
    - calibration intercept (calibration-in-the-large)
    - calibration slope
    - Brier score
with bootstrap 95% confidence intervals.

Expected saved inputs
---------------------
1) Nested CV results from ICI_predict_v2.py:
       <nested_root>/<dataset_stem>/oof_predictions.csv
   Columns expected: id, y_true, p_oof, ...
2) Compare results from Compare validation script:
       either colocated inside each dataset folder under Pre / Post folder
       or under dedicated compare roots such as Pre_Compare / Post_Compare folder
           <root>/<dataset_stem>/nonnested_oof_pred.csv
           <root>/<dataset_stem>/single_split_test_pred.csv
   Columns expected: id, y, p, split

-----
- This is a *supportive internal validation* analysis.
- Path settings are intentionally placed at the top of the script.
"""

from __future__ import annotations
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy.optimize import brentq, minimize
from scipy.special import expit, logit

# ============================================================
# USER SETTINGS
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
RESULTS_DIR = BASE_DIR / "results"
OUTPUT_DIR = RESULTS_DIR / "output"

# Existing result folders
PRE_COMPARE_ROOTS: List[Path] = [
    RESULTS_DIR / "",
    RESULTS_DIR,  # optional fallback if compare results were saved directly under results/<dataset>/
]
POST_COMPARE_ROOTS: List[Path] = [
    RESULTS_DIR / "",
    RESULTS_DIR,
]
PRE_NESTED_ROOT = RESULTS_DIR / ""
POST_NESTED_ROOT = RESULTS_DIR / ""
NMR_NESTED_ROOT = RESULTS_DIR / ""
NMR_COMPARE_ROOTS: List[Path] = [
    RESULTS_DIR / "",          # primary in the current layout
    RESULTS_DIR / "",
    RESULTS_DIR,
]

# Output behavior
CONTINUE_ON_MISSING = True   # If False, stop when any expected file is missing
OVERWRITE_OUTPUT_FILES = True

# Bootstrap / calibration settings
BOOTSTRAP_N = 2000
BOOTSTRAP_SEED = 42
MAX_BOOTSTRAP_ATTEMPTS_FACTOR = 20
MIN_BOOTSTRAP_SUCCESS = None  # auto: max(100, BOOTSTRAP_N // 10)
EPSILON = 1e-6

# Numerical optimization settings
NEWTON_MAX_ITER = 80
NEWTON_TOL = 1e-8
PARAMETER_ABS_BOUND = 25.0

# Plot settings
FIG_DPI = 300
FIG_FORMATS: Tuple[str, ...] = ("tiff", "pdf")
PLOT_ROW_HEIGHT = 0.52
PLOT_WIDTH = 10.8
PLOT_TITLE_FONT = 12
PLOT_LABEL_FONT = 10
PLOT_TICK_FONT = 9
PLOT_MARKER_SIZE = 5
PLOT_CAPSIZE = 2
PLOT_LINEWIDTH = 1.0
PLOT_LEGEND_LOC = "upper center"
SHOW_N_IN_YLABEL = False

# Optional manual x-limits (use None for automatic limits)
SLOPE_XLIM_PRE_NMR: Optional[Tuple[float, float]] = None
INTERCEPT_XLIM_PRE_NMR: Optional[Tuple[float, float]] = None
SLOPE_XLIM_POST: Optional[Tuple[float, float]] = None
INTERCEPT_XLIM_POST: Optional[Tuple[float, float]] = None

# Dataset order and display labels (kept explicit to preserve manuscript order)
DATASET_SPECS: List[Dict[str, str]] = [
    # LC/MS Pre dataset
    {"stem": "input01",      "group": "pre6",   "display": "Dataset 1"},
    {"stem": "input02",       "group": "pre6",   "display": "Dataset 2"},
    {"stem": "input03",      "group": "pre6",   "display": "Dataset 3"},
    {"stem": "input04",    "group": "pre6",   "display": "Dataset 4"},
    {"stem": "",      "group": "pre6",   "display": "Dataset 5"},
    {"stem": "",       "group": "pre6",   "display": "Dataset 6"},

    # LC/MS Post dataset
    {"stem": "",     "group": "post10", "display": "Dataset 1-2"},
    {"stem": "",     "group": "post10", "display": "Dataset 2-2"},
    {"stem": "",    "group": "post10", "display": "Dataset 3-2"},
    {"stem": "",  "group": "post10", "display": "Dataset 4-2"},
    {"stem": "",    "group": "post10", "display": "Dataset 5-2"},
    {"stem": "",      "group": "post10", "display": "Dataset 6-2"},
    {"stem": "",     "group": "post10", "display": "Dataset 2-3"},
    {"stem": "",    "group": "post10", "display": "Dataset 3-3"},
    {"stem": "",  "group": "post10", "display": "Dataset 4-3"},
    {"stem": "",    "group": "post10", "display": "Dataset 5-3"},

    # NMR dataset
    {"stem": "", "group": "nmr1", "display": "Dataset 7 (NMR)"},
]

GROUP_LABELS_EN = {
    "pre6": "Pre-ICI LC/MS",
    "post10": "Post-ICI LC/MS",
    "nmr1": "NMR",
}
GROUP_LABELS_JA = {
    "pre6": "Pre-ICI LC/MS",
    "post10": "Post-ICI LC/MS",
    "nmr1": "NMR",
}

METHOD_ORDER = ["nCV", "CV", "holdout"]
GROUP_ORDER = ["pre6", "post10", "nmr1"]
METRIC_ORDER = ["cal_slope", "cal_intercept", "brier"]

METHOD_FILE_MAP = {
    "nCV": "oof_predictions.csv",
    "CV": "nonnested_oof_pred.csv",
    "holdout": "single_split_test_pred.csv",
}
METHOD_LABELS = {
    "nCV": "nCV",
    "CV": "CV",
    "holdout": "holdout",
}
METHOD_STYLES = {
    "nCV": {"marker": "o", "mfc": "black", "mec": "black", "label": "nCV"},
    "CV": {"marker": "s", "mfc": "white", "mec": "black", "label": "CV"},
    "holdout": {"marker": "^", "mfc": "white", "mec": "black", "label": "holdout"},
}

# ============================================================
# Constants / metadata derived from settings
# ============================================================
DATASET_META: Dict[str, Dict[str, Any]] = {}
for i, spec in enumerate(DATASET_SPECS, start=1):
    DATASET_META[spec["stem"]] = {
        **spec,
        "dataset_order": i,
        "group_order": GROUP_ORDER.index(spec["group"]),
    }

MIN_BOOTSTRAP_SUCCESS_EFFECTIVE = (
    max(100, BOOTSTRAP_N // 10)
    if MIN_BOOTSTRAP_SUCCESS is None
    else int(MIN_BOOTSTRAP_SUCCESS)
)


# ============================================================
# Utilities
# ============================================================
def ensure_output_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)

def to_csv_utf8sig(df: pd.DataFrame, path: Path) -> None:
    df.to_csv(path, index=False, encoding="utf-8-sig")

def fmt_num(x: Any, ndigits: int = 3) -> str:
    try:
        xv = float(x)
    except (TypeError, ValueError):
        return "NA"
    if np.isnan(xv):
        return "NA"
    return f"{xv:.{ndigits}f}"

def fmt_med_range(median_val: Any, min_val: Any, max_val: Any, ndigits: int = 3) -> str:
    return f"{fmt_num(median_val, ndigits)} ({fmt_num(min_val, ndigits)}–{fmt_num(max_val, ndigits)})"

def fmt_med_iqr(median_val: Any, q1_val: Any, q3_val: Any, ndigits: int = 3) -> str:
    return f"{fmt_num(median_val, ndigits)} ({fmt_num(q1_val, ndigits)}–{fmt_num(q3_val, ndigits)})"

def safe_numeric_array(values: Iterable[Any]) -> np.ndarray:
    arr = pd.to_numeric(pd.Series(list(values)), errors="coerce").to_numpy(dtype=float)
    return arr[np.isfinite(arr)]

def add_warning(warnings_rows: List[Dict[str, Any]], dataset: str, method: str, issue: str, detail: str) -> None:
    warnings_rows.append({
        "dataset_stem": dataset,
        "method": method,
        "issue": issue,
        "detail": detail,
    })
    print(f"[WARN] {dataset} | {method} | {issue}: {detail}")

# ============================================================
# File discovery / loading
# ============================================================
def get_candidate_roots(group: str, method: str) -> List[Path]:
    if method == "nCV":
        if group == "pre6":
            return [PRE_NESTED_ROOT]
        if group == "post10":
            return [POST_NESTED_ROOT]
        if group == "nmr1":
            return [NMR_NESTED_ROOT, PRE_NESTED_ROOT]
        raise ValueError(f"Unknown group: {group}")

    if method in {"CV", "holdout"}:
        # In the current project layout, per-dataset compare prediction CSVs are often
        # colocated inside the dataset folders under C.Pre / D.Post rather than under
        # A.Pre_Compare / B.Post_Compare. Search nested roots first, then optional
        # compare roots as fallback.
        if group == "pre6":
            return [PRE_NESTED_ROOT, *PRE_COMPARE_ROOTS]
        if group == "post10":
            return [POST_NESTED_ROOT, *POST_COMPARE_ROOTS]
        if group == "nmr1":
            return [NMR_NESTED_ROOT, PRE_NESTED_ROOT, *NMR_COMPARE_ROOTS]
        raise ValueError(f"Unknown group: {group}")

    raise ValueError(f"Unknown method: {method}")

def find_prediction_file(dataset_stem: str, group: str, method: str, warnings_rows: List[Dict[str, Any]]) -> Optional[Path]:
    filename = METHOD_FILE_MAP[method]
    candidate_roots = get_candidate_roots(group, method)

    exact_candidates: List[Path] = []
    for root in candidate_roots:
        exact_candidates.append(root / dataset_stem / filename)

    for path in exact_candidates:
        if path.exists():
            return path

    recursive_matches: List[Path] = []
    for root in candidate_roots:
        if root.exists():
            recursive_matches.extend(sorted(root.glob(f"**/{dataset_stem}/{filename}")))

    if recursive_matches:
        chosen = sorted(recursive_matches, key=lambda p: (len(p.parts), str(p)))[0]
        if len(recursive_matches) > 1:
            add_warning(
                warnings_rows,
                dataset_stem,
                method,
                "multiple_matches",
                "Multiple matching files found; using the shortest path match: " + str(chosen),
            )
        else:
            add_warning(
                warnings_rows,
                dataset_stem,
                method,
                "recursive_fallback_used",
                f"Exact path not found; using recursive match: {chosen}",
            )
        return chosen

    detail = "Tried: " + " | ".join(str(p) for p in exact_candidates)
    if CONTINUE_ON_MISSING:
        add_warning(warnings_rows, dataset_stem, method, "missing_file", detail)
        return None
    raise FileNotFoundError(detail)


def load_prediction_table(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, encoding="utf-8-sig")

    if {"y_true", "p_oof"}.issubset(df.columns):
        out = pd.DataFrame({
            "id": df.get("id", pd.Series(range(len(df)))),
            "y": pd.to_numeric(df["y_true"], errors="coerce"),
            "p": pd.to_numeric(df["p_oof"], errors="coerce"),
        })
    elif {"y", "p"}.issubset(df.columns):
        out = pd.DataFrame({
            "id": df.get("id", pd.Series(range(len(df)))),
            "y": pd.to_numeric(df["y"], errors="coerce"),
            "p": pd.to_numeric(df["p"], errors="coerce"),
        })
    else:
        raise ValueError(
            f"Unsupported prediction file format: {path}\n"
            f"Columns found: {list(df.columns)}"
        )

    out = out.dropna(subset=["y", "p"]).copy()
    out["y"] = out["y"].astype(int)
    out["p"] = out["p"].astype(float)
    return out.reset_index(drop=True)

# ============================================================
# Calibration metric calculations
# ============================================================
def _neg_log_likelihood(beta: np.ndarray, y: np.ndarray, z: np.ndarray) -> float:
    eta = beta[0] + beta[1] * z
    return float(np.sum(np.logaddexp(0.0, eta) - y * eta))


def _grad_neg_log_likelihood(beta: np.ndarray, y: np.ndarray, z: np.ndarray) -> np.ndarray:
    eta = beta[0] + beta[1] * z
    mu = expit(eta)
    diff = mu - y
    return np.array([np.sum(diff), np.sum(diff * z)], dtype=float)


def fit_calibration_intercept(y: np.ndarray, p: np.ndarray) -> Tuple[float, bool, str]:
    y = np.asarray(y, dtype=float)
    p = np.clip(np.asarray(p, dtype=float), EPSILON, 1.0 - EPSILON)

    if len(np.unique(y)) < 2:
        return np.nan, False, "Only one class in evaluation set; intercept not estimable."

    z = logit(p)
    target = float(np.mean(y))
    if not (0.0 < target < 1.0):
        return np.nan, False, "Event rate is 0 or 1; intercept not finite."

    def f(alpha: float) -> float:
        return float(np.mean(expit(alpha + z)) - target)

    try:
        value = float(brentq(f, -50.0, 50.0, maxiter=200))
        return value, True, ""
    except ValueError as exc:
        return np.nan, False, f"brentq failed for intercept: {exc}"


def fit_calibration_slope(y: np.ndarray, p: np.ndarray) -> Tuple[float, float, bool, str]:
    y = np.asarray(y, dtype=float)
    p = np.clip(np.asarray(p, dtype=float), EPSILON, 1.0 - EPSILON)

    if len(np.unique(y)) < 2:
        return np.nan, np.nan, False, "Only one class in evaluation set; slope not estimable."

    z = logit(p)
    intercept0, ok0, _ = fit_calibration_intercept(y, p)
    if not ok0 or not np.isfinite(intercept0):
        intercept0 = float(logit(np.mean(y)) if 0.0 < np.mean(y) < 1.0 else 0.0)

    beta = np.array([intercept0, 1.0], dtype=float)
    converged = False
    message = ""

    for _ in range(NEWTON_MAX_ITER):
        eta = beta[0] + beta[1] * z
        mu = expit(eta)
        w = np.clip(mu * (1.0 - mu), 1e-12, None)

        score0 = np.sum(y - mu)
        score1 = np.sum((y - mu) * z)
        score = np.array([score0, score1], dtype=float)

        info00 = np.sum(w)
        info01 = np.sum(w * z)
        info11 = np.sum(w * z * z)
        info = np.array([[info00, info01], [info01, info11]], dtype=float)

        det = float(np.linalg.det(info))
        if not np.isfinite(det) or abs(det) < 1e-12:
            message = "Observed information matrix is singular."
            break

        try:
            delta = np.linalg.solve(info, score)
        except np.linalg.LinAlgError:
            message = "Failed to solve Newton step due to singular matrix."
            break

        step = 1.0
        current_nll = _neg_log_likelihood(beta, y, z)
        accepted = False
        while step > 1e-6:
            candidate = beta + step * delta
            if np.max(np.abs(candidate)) > PARAMETER_ABS_BOUND:
                step *= 0.5
                continue
            cand_nll = _neg_log_likelihood(candidate, y, z)
            if np.isfinite(cand_nll) and cand_nll <= current_nll + 1e-12:
                beta = candidate
                accepted = True
                break
            step *= 0.5

        if not accepted:
            message = "Newton step could not be accepted (possible separation / instability)."
            break

        if np.max(np.abs(step * delta)) < NEWTON_TOL:
            converged = True
            message = ""
            break

    if not converged:
        # Fallback to BFGS
        try:
            res = minimize(
                fun=lambda b: _neg_log_likelihood(np.asarray(b, dtype=float), y, z),
                x0=beta,
                jac=lambda b: _grad_neg_log_likelihood(np.asarray(b, dtype=float), y, z),
                method="BFGS",
                options={"gtol": 1e-8, "maxiter": 400},
            )
            if res.success and np.all(np.isfinite(res.x)) and np.max(np.abs(res.x)) <= PARAMETER_ABS_BOUND:
                beta = np.asarray(res.x, dtype=float)
                converged = True
                message = ""
            else:
                message = message or f"BFGS did not converge successfully: {getattr(res, 'message', 'unknown error')}"
        except Exception as exc:  # pragma: no cover - defensive fallback
            message = message or f"BFGS fallback failed: {exc}"

    if not converged:
        return np.nan, np.nan, False, message or "Slope fit did not converge."

    return float(beta[0]), float(beta[1]), True, ""


def compute_point_metrics(y: np.ndarray, p: np.ndarray) -> Dict[str, Any]:
    y = np.asarray(y, dtype=int)
    p = np.asarray(p, dtype=float)
    p = np.clip(p, EPSILON, 1.0 - EPSILON)

    out: Dict[str, Any] = {
        "brier": float(np.mean((p - y) ** 2)),
        "cal_intercept": np.nan,
        "cal_slope": np.nan,
        "fit_ok": True,
        "fit_message": "",
    }

    intercept, ok_i, msg_i = fit_calibration_intercept(y, p)
    out["cal_intercept"] = intercept

    _, slope, ok_s, msg_s = fit_calibration_slope(y, p)
    out["cal_slope"] = slope

    messages = [m for m in [msg_i, msg_s] if m]
    out["fit_ok"] = bool(ok_i and ok_s)
    out["fit_message"] = " | ".join(messages)
    return out


def bootstrap_metric_cis(y: np.ndarray, p: np.ndarray, n_boot: int, seed: int) -> Dict[str, Any]:
    y = np.asarray(y, dtype=int)
    p = np.asarray(p, dtype=float)
    rng = np.random.default_rng(seed)

    intercepts: List[float] = []
    slopes: List[float] = []
    briers: List[float] = []
    attempts = 0
    max_attempts = int(n_boot * MAX_BOOTSTRAP_ATTEMPTS_FACTOR)
    skipped_single_class = 0
    skipped_fit_failure = 0

    while len(slopes) < n_boot and attempts < max_attempts:
        idx = rng.integers(0, len(y), size=len(y))
        yb = y[idx]
        pb = p[idx]
        attempts += 1

        if len(np.unique(yb)) < 2:
            skipped_single_class += 1
            continue

        point = compute_point_metrics(yb, pb)
        if not point["fit_ok"] or not np.isfinite(point["cal_intercept"]) or not np.isfinite(point["cal_slope"]):
            skipped_fit_failure += 1
            continue

        intercepts.append(float(point["cal_intercept"]))
        slopes.append(float(point["cal_slope"]))
        briers.append(float(point["brier"]))

    out: Dict[str, Any] = {
        "n_boot_target": int(n_boot),
        "n_boot_attempts": int(attempts),
        "n_boot_success": int(len(slopes)),
        "n_boot_skipped_single_class": int(skipped_single_class),
        "n_boot_skipped_fit_failure": int(skipped_fit_failure),
        "cal_intercept_ci_lo": np.nan,
        "cal_intercept_ci_hi": np.nan,
        "cal_slope_ci_lo": np.nan,
        "cal_slope_ci_hi": np.nan,
        "brier_ci_lo": np.nan,
        "brier_ci_hi": np.nan,
    }

    if len(slopes) >= MIN_BOOTSTRAP_SUCCESS_EFFECTIVE:
        out["cal_intercept_ci_lo"], out["cal_intercept_ci_hi"] = np.percentile(intercepts, [2.5, 97.5])
        out["cal_slope_ci_lo"], out["cal_slope_ci_hi"] = np.percentile(slopes, [2.5, 97.5])
        out["brier_ci_lo"], out["brier_ci_hi"] = np.percentile(briers, [2.5, 97.5])
    return out

# ============================================================
# Summary tables
# ============================================================
def summarize_group_method(long_df: pd.DataFrame) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []

    for group in GROUP_ORDER:
        for method in METHOD_ORDER:
            sub = long_df[(long_df["group"] == group) & (long_df["method"] == method)].copy()
            if sub.empty:
                rows.append({
                    "group": group,
                    "group_label_en": GROUP_LABELS_EN[group],
                    "group_label_ja": GROUP_LABELS_JA[group],
                    "method": method,
                    "n_datasets": 0,
                    "slope_median": np.nan,
                    "slope_q1": np.nan,
                    "slope_q3": np.nan,
                    "slope_min": np.nan,
                    "slope_max": np.nan,
                    "intercept_median": np.nan,
                    "intercept_q1": np.nan,
                    "intercept_q3": np.nan,
                    "intercept_min": np.nan,
                    "intercept_max": np.nan,
                    "brier_median": np.nan,
                    "brier_q1": np.nan,
                    "brier_q3": np.nan,
                    "brier_min": np.nan,
                    "brier_max": np.nan,
                    "abs_intercept_median": np.nan,
                    "abs_slope_minus1_median": np.nan,
                })
                continue

            slope = safe_numeric_array(sub["cal_slope"].tolist())
            intercept = safe_numeric_array(sub["cal_intercept"].tolist())
            brier = safe_numeric_array(sub["brier"].tolist())

            n_available = int(np.sum(
                np.isfinite(pd.to_numeric(sub["brier"], errors="coerce").to_numpy(dtype=float))
                | np.isfinite(pd.to_numeric(sub["cal_intercept"], errors="coerce").to_numpy(dtype=float))
                | np.isfinite(pd.to_numeric(sub["cal_slope"], errors="coerce").to_numpy(dtype=float))
            ))

            rows.append({
                "group": group,
                "group_label_en": GROUP_LABELS_EN[group],
                "group_label_ja": GROUP_LABELS_JA[group],
                "method": method,
                "n_datasets": n_available,
                "slope_median": float(np.median(slope)) if len(slope) else np.nan,
                "slope_q1": float(np.quantile(slope, 0.25)) if len(slope) else np.nan,
                "slope_q3": float(np.quantile(slope, 0.75)) if len(slope) else np.nan,
                "slope_min": float(np.min(slope)) if len(slope) else np.nan,
                "slope_max": float(np.max(slope)) if len(slope) else np.nan,
                "intercept_median": float(np.median(intercept)) if len(intercept) else np.nan,
                "intercept_q1": float(np.quantile(intercept, 0.25)) if len(intercept) else np.nan,
                "intercept_q3": float(np.quantile(intercept, 0.75)) if len(intercept) else np.nan,
                "intercept_min": float(np.min(intercept)) if len(intercept) else np.nan,
                "intercept_max": float(np.max(intercept)) if len(intercept) else np.nan,
                "brier_median": float(np.median(brier)) if len(brier) else np.nan,
                "brier_q1": float(np.quantile(brier, 0.25)) if len(brier) else np.nan,
                "brier_q3": float(np.quantile(brier, 0.75)) if len(brier) else np.nan,
                "brier_min": float(np.min(brier)) if len(brier) else np.nan,
                "brier_max": float(np.max(brier)) if len(brier) else np.nan,
                "abs_intercept_median": float(np.median(np.abs(intercept))) if len(intercept) else np.nan,
                "abs_slope_minus1_median": float(np.median(np.abs(slope - 1.0))) if len(slope) else np.nan,
            })

    out = pd.DataFrame(rows)
    return out


def build_manuscript_summary(summary_df: pd.DataFrame) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    metrics = [
        ("calibration slope", "slope"),
        ("calibration intercept", "intercept"),
        ("Brier score", "brier"),
    ]

    for group in GROUP_ORDER:
        sdf = summary_df[summary_df["group"] == group].copy().set_index("method")
        for metric_label, prefix in metrics:
            ncv = sdf.loc["nCV"] if "nCV" in sdf.index else None
            cv = sdf.loc["CV"] if "CV" in sdf.index else None
            ho = sdf.loc["holdout"] if "holdout" in sdf.index else None

            ncv_iqr = fmt_med_iqr(ncv[f"{prefix}_median"], ncv[f"{prefix}_q1"], ncv[f"{prefix}_q3"]) if ncv is not None else "NA"
            cv_iqr = fmt_med_iqr(cv[f"{prefix}_median"], cv[f"{prefix}_q1"], cv[f"{prefix}_q3"]) if cv is not None else "NA"
            ho_iqr = fmt_med_iqr(ho[f"{prefix}_median"], ho[f"{prefix}_q1"], ho[f"{prefix}_q3"]) if ho is not None else "NA"

            ncv_range = fmt_med_range(ncv[f"{prefix}_median"], ncv[f"{prefix}_min"], ncv[f"{prefix}_max"]) if ncv is not None else "NA"
            cv_range = fmt_med_range(cv[f"{prefix}_median"], cv[f"{prefix}_min"], cv[f"{prefix}_max"]) if cv is not None else "NA"
            ho_range = fmt_med_range(ho[f"{prefix}_median"], ho[f"{prefix}_min"], ho[f"{prefix}_max"]) if ho is not None else "NA"

            group_en = GROUP_LABELS_EN[group]
            group_ja = GROUP_LABELS_JA[group]
            n_group = int(ncv["n_datasets"]) if ncv is not None and pd.notna(ncv["n_datasets"]) else 0

            text_en = (
                f"In the {group_en} datasets (n={n_group}), {metric_label} was "
                f"median {ncv_iqr} for nCV, {cv_iqr} for CV, and {ho_iqr} for holdout."
            )
            text_ja = (
                f"{group_ja} データセット（n={n_group}）では、{metric_label} は "
                f"nCV で中央値 {ncv_iqr}、CV で {cv_iqr}、holdout で {ho_iqr} であった。"
            )

            rows.append({
                "group": group,
                "group_label_en": group_en,
                "group_label_ja": group_ja,
                "metric": metric_label,
                "nCV": ncv_iqr,
                "CV": cv_iqr,
                "holdout": ho_iqr,
                "nCV_range": ncv_range,
                "CV_range": cv_range,
                "holdout_range": ho_range,
                "text_en": text_en,
                "text_ja": text_ja,
            })

    return pd.DataFrame(rows)


def make_wide_table(long_df: pd.DataFrame) -> pd.DataFrame:
    value_cols = [
        "n_total_dataset",
        "n_eval",
        "n_event",
        "event_rate_eval",
        "mean_pred",
        "brier",
        "brier_ci_lo",
        "brier_ci_hi",
        "cal_intercept",
        "cal_intercept_ci_lo",
        "cal_intercept_ci_hi",
        "cal_slope",
        "cal_slope_ci_lo",
        "cal_slope_ci_hi",
        "n_boot_success",
        "n_boot_attempts",
        "n_boot_skipped_single_class",
        "n_boot_skipped_fit_failure",
        "source_file",
    ]

    parts: List[pd.DataFrame] = []
    base_cols = ["dataset_stem", "display_label", "group", "group_label_en", "group_label_ja", "dataset_order", "group_order"]
    base = long_df[base_cols].drop_duplicates().set_index(["dataset_stem", "display_label", "group", "group_label_en", "group_label_ja", "dataset_order", "group_order"])

    for method in METHOD_ORDER:
        sub = long_df[long_df["method"] == method].copy()
        if sub.empty:
            continue
        sub = sub[["dataset_stem", *value_cols]].set_index("dataset_stem")
        sub = sub.add_suffix(f"_{method}")
        parts.append(sub)

    wide = base
    for part in parts:
        wide = wide.join(part, how="left")

    wide = wide.reset_index().sort_values(["group_order", "dataset_order"]).reset_index(drop=True)
    return wide

# ============================================================
# Plotting
# ============================================================
def _auto_xlim(values: np.ndarray, ref: float, manual_xlim: Optional[Tuple[float, float]]) -> Tuple[float, float]:
    if manual_xlim is not None:
        return manual_xlim

    vals = np.asarray(values, dtype=float)
    vals = vals[np.isfinite(vals)]
    if vals.size == 0:
        return (ref - 1.0, ref + 1.0)

    lo = float(np.min(np.append(vals, ref)))
    hi = float(np.max(np.append(vals, ref)))
    span = max(hi - lo, 0.6)
    pad = max(0.08 * span, 0.05)
    return (lo - pad, hi + pad)


def make_forest_plot(
    df: pd.DataFrame,
    dataset_list: Sequence[str],
    title: str,
    out_base: Path,
    slope_xlim: Optional[Tuple[float, float]],
    intercept_xlim: Optional[Tuple[float, float]],
) -> None:
    plot_df = df[df["dataset_stem"].isin(dataset_list)].copy()
    if plot_df.empty:
        raise ValueError(f"No rows available for forest plot: {title}")

    plot_df["method"] = pd.Categorical(plot_df["method"], categories=METHOD_ORDER, ordered=True)
    order_map = {stem: i for i, stem in enumerate(dataset_list)}
    plot_df["plot_order"] = plot_df["dataset_stem"].map(order_map)
    plot_df = plot_df.sort_values(["plot_order", "method"]).reset_index(drop=True)

    y_base = np.arange(len(dataset_list), dtype=float)
    y_lookup = {stem: y for stem, y in zip(dataset_list, y_base)}
    y_offsets = {"nCV": -0.18, "CV": 0.0, "holdout": 0.18}

    fig_height = max(4.5, 1.8 + PLOT_ROW_HEIGHT * len(dataset_list))
    fig, axes = plt.subplots(1, 2, figsize=(PLOT_WIDTH, fig_height), sharey=True)
    ax_slope, ax_intercept = axes

    handles = []
    labels = []

    slope_extent = []
    intercept_extent = []

    for method in METHOD_ORDER:
        sub = plot_df[plot_df["method"] == method].copy()
        if sub.empty:
            continue

        style = METHOD_STYLES[method]
        ys = np.array([y_lookup[stem] + y_offsets[method] for stem in sub["dataset_stem"]], dtype=float)

        x_slope = sub["cal_slope"].to_numpy(dtype=float)
        x_slope_lo = sub["cal_slope_ci_lo"].to_numpy(dtype=float)
        x_slope_hi = sub["cal_slope_ci_hi"].to_numpy(dtype=float)
        valid_slope = np.isfinite(x_slope) & np.isfinite(x_slope_lo) & np.isfinite(x_slope_hi)
        if np.any(valid_slope):
            h = ax_slope.errorbar(
                x=x_slope[valid_slope],
                y=ys[valid_slope],
                xerr=np.vstack([x_slope[valid_slope] - x_slope_lo[valid_slope], x_slope_hi[valid_slope] - x_slope[valid_slope]]),
                fmt=style["marker"],
                linestyle="none",
                color="black",
                ecolor="black",
                markerfacecolor=style["mfc"],
                markeredgecolor=style["mec"],
                markersize=PLOT_MARKER_SIZE,
                capsize=PLOT_CAPSIZE,
                linewidth=PLOT_LINEWIDTH,
                label=style["label"],
            )
            if method not in labels:
                handles.append(h)
                labels.append(style["label"])
            slope_extent.extend(x_slope_lo[valid_slope].tolist())
            slope_extent.extend(x_slope_hi[valid_slope].tolist())
        else:
            valid_slope = np.isfinite(x_slope)
            if np.any(valid_slope):
                h = ax_slope.plot(
                    x_slope[valid_slope],
                    ys[valid_slope],
                    linestyle="none",
                    marker=style["marker"],
                    color="black",
                    markerfacecolor=style["mfc"],
                    markeredgecolor=style["mec"],
                    markersize=PLOT_MARKER_SIZE,
                    label=style["label"],
                )[0]
                if method not in labels:
                    handles.append(h)
                    labels.append(style["label"])
                slope_extent.extend(x_slope[valid_slope].tolist())

        x_intercept = sub["cal_intercept"].to_numpy(dtype=float)
        x_intercept_lo = sub["cal_intercept_ci_lo"].to_numpy(dtype=float)
        x_intercept_hi = sub["cal_intercept_ci_hi"].to_numpy(dtype=float)
        valid_intercept = np.isfinite(x_intercept) & np.isfinite(x_intercept_lo) & np.isfinite(x_intercept_hi)
        if np.any(valid_intercept):
            ax_intercept.errorbar(
                x=x_intercept[valid_intercept],
                y=ys[valid_intercept],
                xerr=np.vstack([
                    x_intercept[valid_intercept] - x_intercept_lo[valid_intercept],
                    x_intercept_hi[valid_intercept] - x_intercept[valid_intercept],
                ]),
                fmt=style["marker"],
                linestyle="none",
                color="black",
                ecolor="black",
                markerfacecolor=style["mfc"],
                markeredgecolor=style["mec"],
                markersize=PLOT_MARKER_SIZE,
                capsize=PLOT_CAPSIZE,
                linewidth=PLOT_LINEWIDTH,
            )
            intercept_extent.extend(x_intercept_lo[valid_intercept].tolist())
            intercept_extent.extend(x_intercept_hi[valid_intercept].tolist())
        else:
            valid_intercept = np.isfinite(x_intercept)
            if np.any(valid_intercept):
                ax_intercept.plot(
                    x_intercept[valid_intercept],
                    ys[valid_intercept],
                    linestyle="none",
                    marker=style["marker"],
                    color="black",
                    markerfacecolor=style["mfc"],
                    markeredgecolor=style["mec"],
                    markersize=PLOT_MARKER_SIZE,
                )
                intercept_extent.extend(x_intercept[valid_intercept].tolist())

    yticklabels: List[str] = []
    for stem in dataset_list:
        display = DATASET_META[stem]["display"]
        if SHOW_N_IN_YLABEL:
            nvals = plot_df.loc[plot_df["dataset_stem"] == stem, "n_total_dataset"].dropna().unique()
            if len(nvals):
                display = f"{display} (n={int(nvals[0])})"
        yticklabels.append(display)

    for ax in axes:
        ax.set_yticks(y_base)
        ax.set_yticklabels(yticklabels, fontsize=PLOT_TICK_FONT)
        ax.grid(True, axis="x", linestyle=":", linewidth=0.8)
        ax.tick_params(axis="both", labelsize=PLOT_TICK_FONT)

    ax_slope.axvline(1.0, linestyle="--", linewidth=1.0, color="black")
    ax_intercept.axvline(0.0, linestyle="--", linewidth=1.0, color="black")

    ax_slope.set_xlabel("Calibration slope (95% CI)", fontsize=PLOT_LABEL_FONT)
    ax_intercept.set_xlabel("Calibration intercept (95% CI)", fontsize=PLOT_LABEL_FONT)
    ax_slope.set_ylabel("Dataset", fontsize=PLOT_LABEL_FONT)

    ax_slope.set_xlim(_auto_xlim(np.asarray(slope_extent, dtype=float), 1.0, slope_xlim))
    ax_intercept.set_xlim(_auto_xlim(np.asarray(intercept_extent, dtype=float), 0.0, intercept_xlim))

    ax_slope.invert_yaxis()
    fig.suptitle(title, fontsize=PLOT_TITLE_FONT)
    if handles:
        fig.legend(handles, labels, loc=PLOT_LEGEND_LOC, ncol=len(labels), frameon=False, bbox_to_anchor=(0.5, 0.985))

    # subtle separator between pre6 and NMR if present in the combined figure
    if "melanoma-mono_pre" in dataset_list and len(dataset_list) > 1:
        nmr_pos = dataset_list.index("melanoma-mono_pre")
        if nmr_pos > 0:
            sep_y = (y_base[nmr_pos - 1] + y_base[nmr_pos]) / 2.0
            for ax in axes:
                ax.axhline(sep_y, color="black", linewidth=0.7, alpha=0.3)

    fig.tight_layout(rect=[0.0, 0.03, 1.0, 0.95])
    for fmt in FIG_FORMATS:
        fig.savefig(out_base.with_suffix(f".{fmt}"), dpi=FIG_DPI)
    plt.close(fig)

# ============================================================
# Excel / output writing
# ============================================================
def build_readme_sheet() -> pd.DataFrame:
    rows = [
        {"section": "Input files", "item": "nCV", "detail": "Reads oof_predictions.csv saved by ICI_predict_v2.py / ICI_predict.py"},
        {"section": "Input files", "item": "CV", "detail": "Reads nonnested_oof_pred.csv saved by Compare_method_v2.py"},
        {"section": "Input files", "item": "holdout", "detail": "Reads single_split_test_pred.csv saved by Compare_method_v2.py"},
        {"section": "Metric", "item": "Calibration intercept", "detail": "Ideal value = 0. Positive: predictions are globally too low. Negative: predictions are globally too high."},
        {"section": "Metric", "item": "Calibration slope", "detail": "Ideal value = 1. <1: predictions are too extreme. >1: predictions are too moderate."},
        {"section": "Metric", "item": "Brier score", "detail": "Lower is better. Measures overall probabilistic accuracy."},
        {"section": "Interpretation", "item": "Scope", "detail": "Supportive internal analysis only. Not a replacement for external validation."},
        {"section": "Bootstrap", "item": "95% CI", "detail": f"Percentile bootstrap with target {BOOTSTRAP_N} resamples, seed={BOOTSTRAP_SEED}."},
        {"section": "Grouping", "item": "Summary", "detail": "Results are computed dataset-wise, then summarized by Pre-ICI LC/MS, Post-ICI LC/MS, and NMR groups."},
    ]
    return pd.DataFrame(rows)


def build_settings_sheet() -> pd.DataFrame:
    settings_rows: List[Dict[str, Any]] = [
        {"setting": "BASE_DIR", "value": str(BASE_DIR)},
        {"setting": "RESULTS_DIR", "value": str(RESULTS_DIR)},
        {"setting": "OUTPUT_DIR", "value": str(OUTPUT_DIR)},
        {"setting": "PRE_COMPARE_ROOTS", "value": "; ".join(str(p) for p in PRE_COMPARE_ROOTS)},
        {"setting": "POST_COMPARE_ROOTS", "value": "; ".join(str(p) for p in POST_COMPARE_ROOTS)},
        {"setting": "PRE_NESTED_ROOT", "value": str(PRE_NESTED_ROOT)},
        {"setting": "POST_NESTED_ROOT", "value": str(POST_NESTED_ROOT)},
        {"setting": "NMR_NESTED_ROOT", "value": str(NMR_NESTED_ROOT)},
        {"setting": "NMR_COMPARE_ROOTS", "value": "; ".join(str(p) for p in NMR_COMPARE_ROOTS)},
        {"setting": "CONTINUE_ON_MISSING", "value": str(CONTINUE_ON_MISSING)},
        {"setting": "BOOTSTRAP_N", "value": str(BOOTSTRAP_N)},
        {"setting": "BOOTSTRAP_SEED", "value": str(BOOTSTRAP_SEED)},
        {"setting": "MIN_BOOTSTRAP_SUCCESS_EFFECTIVE", "value": str(MIN_BOOTSTRAP_SUCCESS_EFFECTIVE)},
        {"setting": "EPSILON", "value": str(EPSILON)},
        {"setting": "FIG_DPI", "value": str(FIG_DPI)},
        {"setting": "FIG_FORMATS", "value": ", ".join(FIG_FORMATS)},
        {"setting": "created_at", "value": datetime.now().isoformat(timespec="seconds")},
        {"setting": "pandas_version", "value": pd.__version__},
        {"setting": "numpy_version", "value": np.__version__},
    ]
    return pd.DataFrame(settings_rows)


def format_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col_cells in ws.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells[:300]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 48)
    wb.save(path)


def write_outputs(
    long_df: pd.DataFrame,
    wide_df: pd.DataFrame,
    group_summary_df: pd.DataFrame,
    manuscript_df: pd.DataFrame,
    warnings_df: pd.DataFrame,
) -> None:
    ensure_output_dir(OUTPUT_DIR)

    long_csv = OUTPUT_DIR / "calibration_all_long.csv"
    wide_csv = OUTPUT_DIR / "calibration_all_wide.csv"
    group_csv = OUTPUT_DIR / "calibration_group_summary.csv"
    manuscript_csv = OUTPUT_DIR / "calibration_manuscript_summary.csv"
    warnings_csv = OUTPUT_DIR / "warnings.csv"
    settings_json = OUTPUT_DIR / "settings_used.json"
    workbook_path = OUTPUT_DIR / "calibration_support_17datasets.xlsx"

    to_csv_utf8sig(long_df, long_csv)
    to_csv_utf8sig(wide_df, wide_csv)
    to_csv_utf8sig(group_summary_df, group_csv)
    to_csv_utf8sig(manuscript_df, manuscript_csv)
    to_csv_utf8sig(warnings_df, warnings_csv)

    settings_payload = {
        "BASE_DIR": str(BASE_DIR),
        "RESULTS_DIR": str(RESULTS_DIR),
        "OUTPUT_DIR": str(OUTPUT_DIR),
        "PRE_COMPARE_ROOTS": [str(p) for p in PRE_COMPARE_ROOTS],
        "POST_COMPARE_ROOTS": [str(p) for p in POST_COMPARE_ROOTS],
        "PRE_NESTED_ROOT": str(PRE_NESTED_ROOT),
        "POST_NESTED_ROOT": str(POST_NESTED_ROOT),
        "NMR_NESTED_ROOT": str(NMR_NESTED_ROOT),
        "NMR_COMPARE_ROOTS": [str(p) for p in NMR_COMPARE_ROOTS],
        "BOOTSTRAP_N": BOOTSTRAP_N,
        "BOOTSTRAP_SEED": BOOTSTRAP_SEED,
        "MIN_BOOTSTRAP_SUCCESS_EFFECTIVE": MIN_BOOTSTRAP_SUCCESS_EFFECTIVE,
        "EPSILON": EPSILON,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "datasets": DATASET_SPECS,
    }
    settings_json.write_text(json.dumps(settings_payload, indent=2, ensure_ascii=False), encoding="utf-8")

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        build_readme_sheet().to_excel(writer, sheet_name="README", index=False)
        long_df.to_excel(writer, sheet_name="all_long", index=False)
        wide_df.to_excel(writer, sheet_name="all_wide", index=False)
        group_summary_df.to_excel(writer, sheet_name="group_summary", index=False)
        manuscript_df.to_excel(writer, sheet_name="manuscript_summary", index=False)
        warnings_df.to_excel(writer, sheet_name="warnings", index=False)
        build_settings_sheet().to_excel(writer, sheet_name="settings", index=False)

    format_workbook(workbook_path)

# ============================================================
# Main
# ============================================================
def main() -> None:
    ensure_output_dir(OUTPUT_DIR)

    warnings_rows: List[Dict[str, Any]] = []
    rows: List[Dict[str, Any]] = []

    for spec in DATASET_SPECS:
        dataset_stem = spec["stem"]
        group = spec["group"]
        display = spec["display"]
        dataset_order = DATASET_META[dataset_stem]["dataset_order"]
        group_order = DATASET_META[dataset_stem]["group_order"]

        print("=" * 80)
        print(f"[DATASET] {dataset_stem} ({display})")
        print("=" * 80)

        per_method_tables: Dict[str, Optional[pd.DataFrame]] = {}
        per_method_paths: Dict[str, Optional[Path]] = {}
        n_total_candidates: List[int] = []

        for method in METHOD_ORDER:
            path = find_prediction_file(dataset_stem, group, method, warnings_rows)
            per_method_paths[method] = path
            if path is None:
                per_method_tables[method] = None
                continue

            try:
                table = load_prediction_table(path)
            except Exception as exc:
                if CONTINUE_ON_MISSING:
                    add_warning(warnings_rows, dataset_stem, method, "load_failed", str(exc))
                    per_method_tables[method] = None
                    continue
                raise

            if table.empty:
                add_warning(warnings_rows, dataset_stem, method, "empty_table", f"No usable rows after loading {path}")
                per_method_tables[method] = None
                continue

            per_method_tables[method] = table
            n_total_candidates.append(int(len(table)))

        n_total_dataset = max(n_total_candidates) if n_total_candidates else np.nan

        for method in METHOD_ORDER:
            table = per_method_tables.get(method)
            if table is None:
                rows.append({
                    "dataset_stem": dataset_stem,
                    "display_label": display,
                    "group": group,
                    "group_label_en": GROUP_LABELS_EN[group],
                    "group_label_ja": GROUP_LABELS_JA[group],
                    "dataset_order": dataset_order,
                    "group_order": group_order,
                    "method": method,
                    "n_total_dataset": n_total_dataset,
                    "n_eval": np.nan,
                    "n_event": np.nan,
                    "event_rate_eval": np.nan,
                    "mean_pred": np.nan,
                    "brier": np.nan,
                    "brier_ci_lo": np.nan,
                    "brier_ci_hi": np.nan,
                    "cal_intercept": np.nan,
                    "cal_intercept_ci_lo": np.nan,
                    "cal_intercept_ci_hi": np.nan,
                    "cal_slope": np.nan,
                    "cal_slope_ci_lo": np.nan,
                    "cal_slope_ci_hi": np.nan,
                    "abs_intercept": np.nan,
                    "abs_slope_minus1": np.nan,
                    "fit_ok": False,
                    "fit_message": "prediction file missing or could not be loaded",
                    "n_boot_target": BOOTSTRAP_N,
                    "n_boot_attempts": np.nan,
                    "n_boot_success": np.nan,
                    "n_boot_skipped_single_class": np.nan,
                    "n_boot_skipped_fit_failure": np.nan,
                    "source_file": np.nan,
                })
                continue

            y = table["y"].to_numpy(dtype=int)
            p = np.clip(table["p"].to_numpy(dtype=float), EPSILON, 1.0 - EPSILON)

            point = compute_point_metrics(y, p)
            boot = bootstrap_metric_cis(y, p, n_boot=BOOTSTRAP_N, seed=BOOTSTRAP_SEED)

            if not point["fit_ok"]:
                add_warning(warnings_rows, dataset_stem, method, "point_fit_issue", point["fit_message"])
            if int(boot["n_boot_success"]) < MIN_BOOTSTRAP_SUCCESS_EFFECTIVE:
                add_warning(
                    warnings_rows,
                    dataset_stem,
                    method,
                    "bootstrap_low_success",
                    f"Successful bootstrap resamples = {boot['n_boot_success']} (target {BOOTSTRAP_N})",
                )

            row = {
                "dataset_stem": dataset_stem,
                "display_label": display,
                "group": group,
                "group_label_en": GROUP_LABELS_EN[group],
                "group_label_ja": GROUP_LABELS_JA[group],
                "dataset_order": dataset_order,
                "group_order": group_order,
                "method": method,
                "n_total_dataset": n_total_dataset,
                "n_eval": int(len(table)),
                "n_event": int(np.sum(y)),
                "event_rate_eval": float(np.mean(y)),
                "mean_pred": float(np.mean(p)),
                "brier": point["brier"],
                "brier_ci_lo": boot["brier_ci_lo"],
                "brier_ci_hi": boot["brier_ci_hi"],
                "cal_intercept": point["cal_intercept"],
                "cal_intercept_ci_lo": boot["cal_intercept_ci_lo"],
                "cal_intercept_ci_hi": boot["cal_intercept_ci_hi"],
                "cal_slope": point["cal_slope"],
                "cal_slope_ci_lo": boot["cal_slope_ci_lo"],
                "cal_slope_ci_hi": boot["cal_slope_ci_hi"],
                "abs_intercept": abs(point["cal_intercept"]) if np.isfinite(point["cal_intercept"]) else np.nan,
                "abs_slope_minus1": abs(point["cal_slope"] - 1.0) if np.isfinite(point["cal_slope"]) else np.nan,
                "fit_ok": point["fit_ok"],
                "fit_message": point["fit_message"],
                "n_boot_target": boot["n_boot_target"],
                "n_boot_attempts": boot["n_boot_attempts"],
                "n_boot_success": boot["n_boot_success"],
                "n_boot_skipped_single_class": boot["n_boot_skipped_single_class"],
                "n_boot_skipped_fit_failure": boot["n_boot_skipped_fit_failure"],
                "source_file": str(per_method_paths.get(method) or ""),
            }
            rows.append(row)

            print(
                f"[{dataset_stem} | {method}] "
                f"slope={fmt_num(row['cal_slope'])}, intercept={fmt_num(row['cal_intercept'])}, "
                f"Brier={fmt_num(row['brier'])}, bootstrap_success={int(row['n_boot_success'])}"
            )

    long_df = pd.DataFrame(rows)
    long_df["group"] = pd.Categorical(long_df["group"], categories=GROUP_ORDER, ordered=True)
    long_df["method"] = pd.Categorical(long_df["method"], categories=METHOD_ORDER, ordered=True)
    long_df = long_df.sort_values(["group_order", "dataset_order", "method"]).reset_index(drop=True)

    wide_df = make_wide_table(long_df)
    group_summary_df = summarize_group_method(long_df)
    manuscript_df = build_manuscript_summary(group_summary_df)
    warnings_df = pd.DataFrame(warnings_rows, columns=["dataset_stem", "method", "issue", "detail"])

    # CSVs and Excel
    write_outputs(long_df, wide_df, group_summary_df, manuscript_df, warnings_df)

    # Figure data CSVs
    pre_nmr_list = [d["stem"] for d in DATASET_SPECS if d["group"] in {"pre6", "nmr1"}]
    post_list = [d["stem"] for d in DATASET_SPECS if d["group"] == "post10"]

    pre_nmr_df = long_df[long_df["dataset_stem"].isin(pre_nmr_list)].copy()
    post_df = long_df[long_df["dataset_stem"].isin(post_list)].copy()
    to_csv_utf8sig(pre_nmr_df, OUTPUT_DIR / "figuredata_pre6_nmr_long.csv")
    to_csv_utf8sig(post_df, OUTPUT_DIR / "figuredata_post10_long.csv")

    # Forest plots
    make_forest_plot(
        df=long_df,
        dataset_list=pre_nmr_list,
        title="",
        out_base=OUTPUT_DIR / "calibration_forest_pre6_nmr",
        slope_xlim=SLOPE_XLIM_PRE_NMR,
        intercept_xlim=INTERCEPT_XLIM_PRE_NMR,
    )
    make_forest_plot(
        df=long_df,
        dataset_list=post_list,
        title="",
        out_base=OUTPUT_DIR / "calibration_forest_post10",
        slope_xlim=SLOPE_XLIM_POST,
        intercept_xlim=INTERCEPT_XLIM_POST,
    )

    print("\n[OK] Saved outputs to:")
    print(f"- {OUTPUT_DIR}")
    print(f"- {OUTPUT_DIR / 'calibration_support_17datasets.xlsx'}")
    print(f"- {OUTPUT_DIR / 'calibration_forest_pre6_nmr.pdf'}")
    print(f"- {OUTPUT_DIR / 'calibration_forest_post10.pdf'}")

if __name__ == "__main__":
    main()