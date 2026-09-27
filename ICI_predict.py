#!/usr/bin/env python3
"""
Leakage-aware metabolomics classification with repeated nested CV.

This implementation uses the established model sequence:

    Elastic-net tuning -> resampling-based feature ranking ->
    Top-k/L2 tuning by the 1-SE rule -> outer-fold prediction

The implementation adds strict input QC, scenario-based sensitivity analyses,
repeat/fold-level audit outputs, feature-stability source data, and automatic
CV fold reduction when the minority class is smaller than the requested fold
count. Threshold-dependent metrics, Youden thresholds, confusion matrices,
and decision-curve analysis are intentionally not produced.

Standard input names
--------------------
    dataset1_pre.xlsx, dataset1_post1.xlsx, dataset1_post2.xlsx, ...

Standard sheet contract
-----------------------
    Sheet1: primary binary analysis
    Sheet2: Dataset 1, SD classified as non-responder
    Sheet3: Dataset 1, SD classified as responder

Each analysis sheet must contain:
    column 1: id
    column 2: category or Responder
    remaining columns: numeric molecular features

Primary outputs retain compatibility aliases used by the existing downstream
scripts: metrics_summary.csv, tuning_summary.csv, and oof_predictions.csv.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
import re
import sys
import time
import warnings
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from itertools import combinations
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Mapping, Optional, Sequence, Tuple, Union

import numpy as np
import openpyxl
import pandas as pd
import sklearn
from sklearn.exceptions import ConvergenceWarning
from sklearn.impute import KNNImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    average_precision_score,
    brier_score_loss,
    precision_recall_curve,
    roc_auc_score,
    roc_curve,
)
from sklearn.model_selection import RepeatedStratifiedKFold

import dataset_selection
from dataset_selection import (
    FORMAL_PRE_KEYS, KNOWN_DATASET_KEYS, canonical_key, coverage_rows, select_available,
)


# =============================================================================
# USER SETTINGS: edit this section
# =============================================================================

SCRIPT_DIR = Path(__file__).resolve().parent
INPUT_DIR = SCRIPT_DIR / "raw"
OUTPUT_ROOT = SCRIPT_DIR / "results_revision"
RUN_TAG = "All_data_primary"

# If empty, files are discovered recursively under INPUT_DIR using the strict
# datasetN_pre/post1/post2.xlsx naming rule. If populated, standard filenames
# are still required unless an entry is also defined in INPUT_METADATA_OVERRIDES.
INPUT_FILES: List[str] = []
SEARCH_INPUT_RECURSIVELY = True

# Optional metadata for a non-standard filename. Keys must be absolute paths or
# paths resolvable from the current working directory.
# Example:
# INPUT_METADATA_OVERRIDES = {
#     r"/path/Dataset4_Pre.xlsx": {
#         "dataset_id": "dataset4", "timepoint": "pre", "sheet_name": "Analysis"
#     }
# }
INPUT_METADATA_OVERRIDES: Dict[str, Dict[str, str]] = {}

# Optional primary-sheet overrides for already prepared workbooks whose primary
# sheet is not Sheet1 (for example, historical files with an Analysis sheet).
PRIMARY_SHEET_OVERRIDES: Dict[str, str] = {}

ID_COL = "id"
LABEL_COLUMN_CANDIDATES: Tuple[str, ...] = ("category", "Responder")
PRIMARY_SHEET = "Sheet1"
SD_AS_NR_SHEET = "Sheet2"
SD_AS_R_SHEET = "Sheet3"

# Restrict the entire invocation to selected canonical keys, e.g.
# ["dataset2_pre", "dataset6_pre"]. Use "ALL" for every discovered input.
DATASETS_TO_RUN = "ALL"
PRIMARY_TARGETS= "ALL" #: Union[str, List[str]] = "ALL"
SD_TARGETS: List[str] = ["dataset1_pre", "dataset1_post1"]
KNN_TARGETS: List[str] = list(FORMAL_PRE_KEYS)
OUTER3_TARGETS: Union[str, List[str]] = "ALL_PRIMARY"

RUN_PRIMARY = True
RUN_SD_SENSITIVITY = False
RUN_KNN_SENSITIVITY = False
RUN_OUTER3_SENSITIVITY = False

# When True, a kmax30 analysis is scheduled for any primary analysis in the
# current invocation where at least one outer fit selected the primary ceiling.
RUN_KMAX_SENSITIVITY = False
AUTO_KMAX_FROM_PRIMARY = True
KMAX_EXPLICIT_TARGETS: List[str] = []

RANDOM_STATE = 42

# Main repeated nested-CV design.
OUTER_SPLITS_TARGET = 5
OUTER_REPEATS = 5
INNER_SPLITS_TARGET = 3
INNER_REPEATS = 5
OUTER3_SPLITS_TARGET = 3

MISSING_RATE_THRESHOLD = 0.50
PRIMARY_MAX_K = 10
KMAX_SENSITIVITY_MAX_K = 30
KNN_NEIGHBORS = 5

# Compatibility alias used by existing downstream scripts.
MAX_K = PRIMARY_MAX_K

EN_C_GRID = [1e-3, 3e-3, 1e-2, 3e-2, 1e-1, 1, 3]
EN_L1_RATIO_GRID = [0, 0.05, 0.2, 0.5, 0.8, 0.95]
L2_C_GRID = [1e-4, 3e-4, 1e-3, 3e-3, 1e-2, 3e-2, 1e-1, 3e-1, 1, 3]

EN_MAX_ITER_INIT = 5000
EN_MAX_ITER_LIMIT = 20000
L2_MAX_ITER_INIT = 2000
L2_MAX_ITER_LIMIT = 10000
TOL = 1e-4
HARD_FAIL_ON_NONCONVERGENCE = True

# A compatibility/descriptive patient bootstrap. Because repeated-CV OOF
# predictions are dependent, this CI is not treated as CV-variance corrected.
BOOTSTRAP_N = 2000

# Existing result folders are never recursively deleted. Set this to True only
# when intentionally replacing files inside an already separated run folder.
ALLOW_OVERWRITE = False

# Known technical/internal-standard features. Presence causes a fail-fast QC
# error so an upstream analysis matrix cannot silently reintroduce them.
FORBIDDEN_EXACT_BY_DATASET: Dict[str, Tuple[str, ...]] = {
    "dataset1": ("L-Met",),
    "dataset2": ("TIC", "ALL1", "L-Met", "(-)-Isoproterenol"),
}
FORBIDDEN_REGEX_BY_DATASET: Dict[str, Tuple[str, ...]] = {
    "dataset2": (r"(?i)^13c[23](?:\b|[_-])",),
}

# Feature-specific preprocessing exceptions. All other features retain the
# historical positive-intensity policy: negative/zero -> missing, half-minimum
# imputation, log1p, then z-score within the training fold.
FEATURE_RULE_OVERRIDES: Dict[str, Dict[str, Dict[str, Any]]] = {
    "dataset4": {
        "Kyn/Trp (log10)": {
            "transform": "identity",
            "zero_is_missing": False,
            "negative_is_missing": False,
            "half_minimum_allowed": False,
        }
    }
}
REQUIRED_FEATURES_BY_DATASET: Dict[str, Tuple[str, ...]] = {
    "dataset4": ("Kyn/Trp (log10)",),
}


# =============================================================================
# PUBLIC SCENARIO PROFILES
# =============================================================================

# The original analysis used one copy of this script per sensitivity analysis.
# The public version keeps one implementation and applies only the setting
# differences recorded in those copies. Calling ``configure_scenario`` is
# deterministic: every setting below is reset from a fixed profile each time.
PUBLIC_SCENARIOS: Tuple[str, ...] = (
    "primary",
    "sd",
    "knn",
    "kmax30",
    "outer3",
    "reggrid",
)

_ALL_PRIMARY_KEYS: Tuple[str, ...] = (
    "dataset1_pre", "dataset1_post1",
    "dataset2_pre", "dataset2_post1", "dataset2_post2",
    "dataset3_pre", "dataset3_post1", "dataset3_post2",
    "dataset4_pre", "dataset4_post1", "dataset4_post2",
    "dataset5_pre", "dataset5_post1", "dataset5_post2",
    "dataset6_pre", "dataset6_post1",
    "dataset7_pre",
)

_REGGRID_KEYS: Tuple[str, ...] = (
    "dataset1_pre",
    "dataset2_pre",
    "dataset3_pre",
    "dataset4_pre",
    "dataset5_pre",
    "dataset7_pre",
)

_PRIMARY_EN_C_GRID: Tuple[float, ...] = (1e-3, 3e-3, 1e-2, 3e-2, 1e-1, 1, 3)
_PRIMARY_L2_C_GRID: Tuple[float, ...] = (
    1e-4, 3e-4, 1e-3, 3e-3, 1e-2, 3e-2, 1e-1, 3e-1, 1, 3,
)
_REGGRID_EN_C_GRID: Tuple[float, ...] = (
    1e-4, 3e-4, 1e-3, 3e-3, 1e-2, 3e-2, 1e-1, 1, 3, 10, 30,
)
_REGGRID_L2_C_GRID: Tuple[float, ...] = (
    1e-5, 3e-5, 1e-4, 3e-4, 1e-3, 3e-3, 1e-2,
    3e-2, 1e-1, 3e-1, 1, 3, 10, 30,
)

_SCENARIO_PROFILES: Dict[str, Dict[str, Any]] = {
    "primary": {
        "description": "Primary repeated nested cross-validation analyses.",
        "run_tag": "All_data_primary",
        "datasets_to_run": "ALL",
        "primary_targets": "ALL",
        "sd_targets": ("dataset1_pre", "dataset1_post1"),
        "knn_targets": FORMAL_PRE_KEYS,
        "outer3_targets": "ALL_PRIMARY",
        "run_primary": True,
        "run_sd": False,
        "run_knn": False,
        "run_outer3": False,
        "run_kmax": False,
        "auto_kmax": True,
        "kmax_targets": (),
        "en_C_grid": _PRIMARY_EN_C_GRID,
        "l2_C_grid": _PRIMARY_L2_C_GRID,
        "hard_fail_on_nonconvergence": True,
        "primary_scenario_name": "primary",
        "primary_scenario_source": "primary",
    },
    "sd": {
        "description": "Stable disease classified as non-responder or responder.",
        "run_tag": "All_data_SD",
        "datasets_to_run": ("dataset1_pre", "dataset1_post1"),
        "primary_targets": "ALL",
        "sd_targets": ("dataset1_pre", "dataset1_post1"),
        "knn_targets": FORMAL_PRE_KEYS,
        "outer3_targets": "ALL_PRIMARY",
        "run_primary": False,
        "run_sd": True,
        "run_knn": False,
        "run_outer3": False,
        "run_kmax": False,
        "auto_kmax": True,
        "kmax_targets": (),
        "en_C_grid": _PRIMARY_EN_C_GRID,
        "l2_C_grid": _PRIMARY_L2_C_GRID,
        "hard_fail_on_nonconvergence": True,
        "primary_scenario_name": "primary",
        "primary_scenario_source": "primary",
    },
    "knn": {
        "description": "KNN preprocessing sensitivity in the six formal pre-ICI datasets.",
        "run_tag": "All_data_KNN",
        "datasets_to_run": FORMAL_PRE_KEYS,
        "primary_targets": "ALL",
        "sd_targets": ("dataset2_pre",),
        "knn_targets": FORMAL_PRE_KEYS,
        "outer3_targets": "ALL_PRIMARY",
        "run_primary": False,
        "run_sd": False,
        "run_knn": True,
        "run_outer3": False,
        "run_kmax": False,
        "auto_kmax": True,
        "kmax_targets": (),
        "en_C_grid": _PRIMARY_EN_C_GRID,
        "l2_C_grid": _PRIMARY_L2_C_GRID,
        "hard_fail_on_nonconvergence": True,
        "primary_scenario_name": "primary",
        "primary_scenario_source": "primary",
    },
    "kmax30": {
        "description": "Maximum top-k increased from 10 to 30.",
        "run_tag": "All_data_Kmax30",
        "datasets_to_run": "ALL",
        "primary_targets": "ALL",
        "sd_targets": ("dataset1_pre", "dataset1_post1"),
        "knn_targets": FORMAL_PRE_KEYS,
        "outer3_targets": "ALL_PRIMARY",
        "run_primary": False,
        "run_sd": False,
        "run_knn": False,
        "run_outer3": False,
        "run_kmax": True,
        "auto_kmax": False,
        "kmax_targets": _ALL_PRIMARY_KEYS,
        "en_C_grid": _PRIMARY_EN_C_GRID,
        "l2_C_grid": _PRIMARY_L2_C_GRID,
        "hard_fail_on_nonconvergence": True,
        "primary_scenario_name": "primary",
        "primary_scenario_source": "primary",
    },
    "outer3": {
        "description": "Three-fold outer cross-validation sensitivity analysis.",
        "run_tag": "All_data_outer3",
        "datasets_to_run": "ALL",
        "primary_targets": "ALL",
        "sd_targets": ("dataset1_pre", "dataset1_post1"),
        "knn_targets": FORMAL_PRE_KEYS,
        "outer3_targets": "ALL",
        "run_primary": False,
        "run_sd": False,
        "run_knn": False,
        "run_outer3": True,
        "run_kmax": False,
        "auto_kmax": True,
        "kmax_targets": (),
        "en_C_grid": _PRIMARY_EN_C_GRID,
        "l2_C_grid": _PRIMARY_L2_C_GRID,
        "hard_fail_on_nonconvergence": True,
        "primary_scenario_name": "primary",
        "primary_scenario_source": "primary",
    },
    "reggrid": {
        "description": "Expanded regularization-grid sensitivity analysis.",
        "run_tag": "All_data_RegGrid",
        "datasets_to_run": _REGGRID_KEYS,
        "primary_targets": _REGGRID_KEYS,
        "sd_targets": ("dataset1_pre", "dataset1_post1"),
        "knn_targets": FORMAL_PRE_KEYS,
        "outer3_targets": "ALL_PRIMARY",
        "run_primary": True,
        "run_sd": False,
        "run_knn": False,
        "run_outer3": False,
        "run_kmax": False,
        "auto_kmax": True,
        "kmax_targets": (),
        "en_C_grid": _REGGRID_EN_C_GRID,
        "l2_C_grid": _REGGRID_L2_C_GRID,
        "hard_fail_on_nonconvergence": False,
        "primary_scenario_name": "reggrid",
        "primary_scenario_source": "regularization_grid_sensitivity",
    },
}

ACTIVE_SCENARIO = "primary"
ACTIVE_SCENARIO_METADATA: Dict[str, Any] = {}
# A CLI --datasets request is always strict, including when equal to a profile.
EXPLICIT_DATASET_SELECTION = False


def _selector_copy(value: Any) -> Any:
    """Return a mutable list for stored selectors while preserving sentinels."""
    return value if isinstance(value, str) else list(value)


def get_public_scenario_profile(name: str) -> Dict[str, Any]:
    """Return a detached, JSON-friendly profile without changing runtime state."""
    if name not in _SCENARIO_PROFILES:
        raise ValueError(f"Unknown scenario {name!r}; choose from {PUBLIC_SCENARIOS}")
    profile = _SCENARIO_PROFILES[name]
    return {
        key: _selector_copy(value) if isinstance(value, (str, tuple)) else value
        for key, value in profile.items()
    }


def configure_scenario(name: str) -> Dict[str, Any]:
    """Apply one recorded scenario profile and return its detached metadata."""
    global ACTIVE_SCENARIO, ACTIVE_SCENARIO_METADATA, EXPLICIT_DATASET_SELECTION
    global RUN_TAG, DATASETS_TO_RUN, PRIMARY_TARGETS, SD_TARGETS, KNN_TARGETS
    global OUTER3_TARGETS, RUN_PRIMARY, RUN_SD_SENSITIVITY
    global RUN_KNN_SENSITIVITY, RUN_OUTER3_SENSITIVITY
    global RUN_KMAX_SENSITIVITY, AUTO_KMAX_FROM_PRIMARY, KMAX_EXPLICIT_TARGETS
    global EN_C_GRID, L2_C_GRID, HARD_FAIL_ON_NONCONVERGENCE, MAX_K

    profile = get_public_scenario_profile(name)
    ACTIVE_SCENARIO = name
    EXPLICIT_DATASET_SELECTION = False
    RUN_TAG = str(profile["run_tag"])
    DATASETS_TO_RUN = _selector_copy(profile["datasets_to_run"])
    PRIMARY_TARGETS = _selector_copy(profile["primary_targets"])
    SD_TARGETS = _selector_copy(profile["sd_targets"])
    KNN_TARGETS = _selector_copy(profile["knn_targets"])
    OUTER3_TARGETS = _selector_copy(profile["outer3_targets"])
    RUN_PRIMARY = bool(profile["run_primary"])
    RUN_SD_SENSITIVITY = bool(profile["run_sd"])
    RUN_KNN_SENSITIVITY = bool(profile["run_knn"])
    RUN_OUTER3_SENSITIVITY = bool(profile["run_outer3"])
    RUN_KMAX_SENSITIVITY = bool(profile["run_kmax"])
    AUTO_KMAX_FROM_PRIMARY = bool(profile["auto_kmax"])
    KMAX_EXPLICIT_TARGETS = _selector_copy(profile["kmax_targets"])
    EN_C_GRID = list(profile["en_C_grid"])
    L2_C_GRID = list(profile["l2_C_grid"])
    HARD_FAIL_ON_NONCONVERGENCE = bool(profile["hard_fail_on_nonconvergence"])
    MAX_K = PRIMARY_MAX_K
    ACTIVE_SCENARIO_METADATA = profile
    return get_public_scenario_profile(name)


# Imports must always expose the primary-analysis defaults. CLI selection is
# intentionally performed only in the ``__main__`` block below.
configure_scenario("primary")


# =============================================================================
# DATA STRUCTURES
# =============================================================================

TARGET_SELECTOR = Union[str, Sequence[str]]
STANDARD_NAME_RE = re.compile(
    r"^dataset(?P<number>[1-7])_(?P<timepoint>pre|post1|post2)\.xlsx$",
    flags=re.IGNORECASE,
)


@dataclass(frozen=True)
class FeatureRule:
    transform: str = "log1p"
    zero_is_missing: bool = True
    negative_is_missing: bool = True
    half_minimum_allowed: bool = True

    def validate(self) -> None:
        if self.transform not in {"log1p", "identity"}:
            raise ValueError(f"Unsupported feature transform: {self.transform}")


@dataclass(frozen=True)
class DatasetSpec:
    path: Path
    dataset_id: str
    timepoint: str
    key: str
    primary_sheet: str = PRIMARY_SHEET

    @property
    def dataset_number(self) -> int:
        return int(self.dataset_id.removeprefix("dataset"))

    @property
    def primary_role(self) -> str:
        if self.dataset_id == "dataset6":
            return "supplement_small_sample"
        if self.dataset_id == "dataset7":
            return "nmr_separate"
        return "main_lcms"


@dataclass(frozen=True)
class ScenarioSpec:
    name: str
    sheet_name: str
    imputation: str = "half_minimum"
    max_k: int = PRIMARY_MAX_K
    outer_splits_target: int = OUTER_SPLITS_TARGET
    source: str = "primary"

    def validate(self) -> None:
        if self.imputation not in {"half_minimum", "knn"}:
            raise ValueError(f"Unsupported imputation strategy: {self.imputation}")
        if self.max_k < 1:
            raise ValueError("max_k must be >= 1")
        if self.outer_splits_target < 2:
            raise ValueError("outer_splits_target must be >= 2")


@dataclass(frozen=True)
class SplitRecord:
    repeat: int
    fold: int
    train_idx: np.ndarray
    test_idx: np.ndarray


class CVPlan:
    """Materialized repeated stratified CV with stable repeat/fold labels."""

    def __init__(self, records: Sequence[SplitRecord], n_splits: int, n_repeats: int):
        self.records = list(records)
        self.n_splits = int(n_splits)
        self.n_repeats = int(n_repeats)

    def split(self, X: Any = None, y: Any = None, groups: Any = None) -> Iterator[Tuple[np.ndarray, np.ndarray]]:
        del X, y, groups
        for record in self.records:
            yield record.train_idx, record.test_idx

    def get_n_splits(self, X: Any = None, y: Any = None, groups: Any = None) -> int:
        del X, y, groups
        return len(self.records)


@dataclass
class DataBundle:
    spec: DatasetSpec
    sheet_name: str
    X: pd.DataFrame
    y: np.ndarray
    ids: pd.Series
    feature_names: List[str]
    feature_rules: Dict[str, FeatureRule]
    feature_manifest: pd.DataFrame
    label_col: str
    input_sha256: str
    input_warnings: List[str] = field(default_factory=list)


@dataclass
class RunResult:
    dataset_key: str
    scenario: str
    output_dir: Path
    outer_splits_used: int
    max_selected_k: int
    primary_k_ceiling_hit: bool
    metrics: Dict[str, Any]


# =============================================================================
# GENERAL UTILITIES
# =============================================================================

def _normalize_name(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def _json_default(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, FeatureRule):
        return asdict(value)
    raise TypeError(f"Object of type {type(value).__name__} is not JSON serializable")


def write_json(path: Path, payload: Mapping[str, Any]) -> None:
    path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=_json_default) + "\n",
        encoding="utf-8",
    )


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def to_csv_utf8sig(df: pd.DataFrame, path: Path, index: bool = False) -> None:
    df.to_csv(path, index=index, encoding="utf-8-sig")


def to_tsv_utf8sig(df: pd.DataFrame, path: Path, index: bool = False) -> None:
    df.to_csv(path, sep="\t", index=index, encoding="utf-8-sig")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def prepare_output_dir(path: Path) -> None:
    if path.exists() and any(path.iterdir()) and not ALLOW_OVERWRITE:
        raise FileExistsError(
            f"Output directory already contains files: {path}\n"
            "Change RUN_TAG/OUTPUT_ROOT or set ALLOW_OVERWRITE=True intentionally."
        )
    path.mkdir(parents=True, exist_ok=True)


def class_counts(y: np.ndarray) -> Dict[str, int]:
    values = np.asarray(y, dtype=int)
    return {
        "n": int(len(values)),
        "positive": int(np.sum(values == 1)),
        "negative": int(np.sum(values == 0)),
    }


def adjusted_n_splits(
    y: np.ndarray,
    target_splits: int,
    context: str,
    log_lines: List[str],
) -> int:
    values = np.asarray(y, dtype=int)
    classes, counts = np.unique(values, return_counts=True)
    if set(classes.tolist()) != {0, 1}:
        raise ValueError(f"{context}: both binary classes are required; found {classes.tolist()}")
    minimum = int(counts.min())
    if minimum < 2:
        raise ValueError(f"{context}: minority class count is {minimum}; CV requires at least 2")
    used = min(int(target_splits), minimum)
    level = "WARN" if used < target_splits else "INFO"
    message = (
        f"[{level}] {context}: requested_splits={target_splits}, "
        f"used_splits={used}, minority_class_n={minimum}"
    )
    print(message)
    log_lines.append(message)
    return used


def make_cv_plan(
    y: np.ndarray,
    target_splits: int,
    n_repeats: int,
    context: str,
    log_lines: List[str],
    random_state: int = RANDOM_STATE,
) -> CVPlan:
    n_splits = adjusted_n_splits(y, target_splits, context, log_lines)
    cv = RepeatedStratifiedKFold(
        n_splits=n_splits,
        n_repeats=int(n_repeats),
        random_state=int(random_state),
    )
    dummy = np.zeros((len(y), 1), dtype=float)
    records: List[SplitRecord] = []
    for index, (train_idx, test_idx) in enumerate(cv.split(dummy, y)):
        records.append(
            SplitRecord(
                repeat=index // n_splits + 1,
                fold=index % n_splits + 1,
                train_idx=np.asarray(train_idx, dtype=int),
                test_idx=np.asarray(test_idx, dtype=int),
            )
        )
    return CVPlan(records, n_splits=n_splits, n_repeats=n_repeats)


def make_inner_cv(y_train: np.ndarray, context: str, log_lines: List[str]) -> CVPlan:
    """Compatibility entry point used by downstream scripts."""
    return make_cv_plan(
        y_train,
        target_splits=INNER_SPLITS_TARGET,
        n_repeats=INNER_REPEATS,
        context=context,
        log_lines=log_lines,
        random_state=RANDOM_STATE,
    )


def _canonical_target(value: str) -> str:
    return str(value).strip().casefold()


def select_targets(selector: TARGET_SELECTOR, available: Iterable[str]) -> List[str]:
    """Strict explicit selection; ALL selects the supplied known inputs."""
    return select_available(selector, available)


def _profile_targets(
    profile_key: str, selector: TARGET_SELECTOR, available: Iterable[str],
) -> List[str]:
    """Profile defaults intersect availability; user overrides remain strict."""
    default = _SCENARIO_PROFILES[ACTIVE_SCENARIO][profile_key]
    normalized = selector if isinstance(selector, str) else tuple(selector)
    if normalized == default and not (
        profile_key == "datasets_to_run" and EXPLICIT_DATASET_SELECTION
    ):
        default_keys = None if isinstance(default, str) else default
        return select_available("ALL", available, default_keys=default_keys)
    return select_targets(selector, available)


def _resolve_user_path(value: str) -> Path:
    path = Path(value).expanduser()
    return path.resolve() if path.is_absolute() else (Path.cwd() / path).resolve()


def discover_datasets() -> Dict[str, DatasetSpec]:
    candidates: List[Path]
    if INPUT_FILES:
        candidates = [_resolve_user_path(item) for item in INPUT_FILES]
    else:
        if not INPUT_DIR.exists():
            raise FileNotFoundError(f"INPUT_DIR does not exist: {INPUT_DIR}")
        iterator = INPUT_DIR.rglob("*.xlsx") if SEARCH_INPUT_RECURSIVELY else INPUT_DIR.glob("*.xlsx")
        candidates = sorted(path.resolve() for path in iterator if not path.name.startswith("~$"))

    override_lookup = {_resolve_user_path(path): dict(meta) for path, meta in INPUT_METADATA_OVERRIDES.items()}
    specs: Dict[str, DatasetSpec] = {}

    for path in candidates:
        if not path.exists():
            raise FileNotFoundError(f"Input file does not exist: {path}")

        override = override_lookup.get(path)
        match = STANDARD_NAME_RE.match(path.name)
        if override is None and match is None:
            raise ValueError(
                f"Non-standard input filename: {path.name}. "
                "Keep only canonical study inputs under INPUT_DIR, or specify "
                "INPUT_FILES/INPUT_METADATA_OVERRIDES explicitly."
            )

        if override is not None:
            dataset_id = _canonical_target(override["dataset_id"])
            timepoint = _canonical_target(override["timepoint"])
            primary_sheet = override.get("sheet_name", PRIMARY_SHEET)
        else:
            assert match is not None
            dataset_id = f"dataset{int(match.group('number'))}"
            timepoint = match.group("timepoint").casefold()
            primary_sheet = PRIMARY_SHEET_OVERRIDES.get(f"{dataset_id}_{timepoint}", PRIMARY_SHEET)

        if not re.fullmatch(r"dataset[1-7]", dataset_id):
            raise ValueError(f"Invalid dataset_id for {path}: {dataset_id}")
        if timepoint not in {"pre", "post1", "post2"}:
            raise ValueError(f"Invalid timepoint for {path}: {timepoint}")

        key = canonical_key(f"{dataset_id}_{timepoint}")
        if key in specs:
            raise ValueError(f"Multiple input files resolve to {key}: {specs[key].path} and {path}")
        specs[key] = DatasetSpec(
            path=path,
            dataset_id=dataset_id,
            timepoint=timepoint,
            key=key,
            primary_sheet=primary_sheet,
        )

    if not specs:
        raise RuntimeError("No valid datasetN_pre/post1/post2.xlsx inputs were found")
    return specs


# =============================================================================
# INPUT QC AND FEATURE RULES
# =============================================================================

def _default_feature_rule(feature: str) -> FeatureRule:
    if _normalize_name(feature) == _normalize_name("Kyn/Trp (log10)"):
        return FeatureRule(
            transform="identity",
            zero_is_missing=False,
            negative_is_missing=False,
            half_minimum_allowed=False,
        )
    return FeatureRule()


def resolve_feature_rules(feature_names: Sequence[str], dataset_id: Optional[str]) -> Dict[str, FeatureRule]:
    rules = {feature: _default_feature_rule(feature) for feature in feature_names}
    if not dataset_id:
        return rules

    normalized_actual = {_normalize_name(feature): feature for feature in feature_names}
    overrides = FEATURE_RULE_OVERRIDES.get(dataset_id, {})
    for requested_name, values in overrides.items():
        actual = normalized_actual.get(_normalize_name(requested_name))
        if actual is None:
            continue
        rule = FeatureRule(**values)
        rule.validate()
        rules[actual] = rule
    return rules


def _validate_required_and_forbidden_features(feature_names: Sequence[str], dataset_id: str) -> None:
    normalized = {_normalize_name(feature): feature for feature in feature_names}
    missing_required = [
        feature
        for feature in REQUIRED_FEATURES_BY_DATASET.get(dataset_id, ())
        if _normalize_name(feature) not in normalized
    ]
    if missing_required:
        raise ValueError(f"{dataset_id}: required features are missing: {missing_required}")

    forbidden: List[str] = []
    for requested in FORBIDDEN_EXACT_BY_DATASET.get(dataset_id, ()):
        actual = normalized.get(_normalize_name(requested))
        if actual is not None:
            forbidden.append(actual)
    for pattern in FORBIDDEN_REGEX_BY_DATASET.get(dataset_id, ()):
        regex = re.compile(pattern)
        forbidden.extend(feature for feature in feature_names if regex.search(str(feature)))
    if forbidden:
        raise ValueError(
            f"{dataset_id}: forbidden technical/internal-standard features were found: "
            f"{sorted(set(forbidden))}"
        )


def _read_raw_headers(path: Path, sheet_name: str) -> List[str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Required sheet '{sheet_name}' is missing in {path.name}; "
                f"available={workbook.sheetnames}"
            )
        worksheet = workbook[sheet_name]
        raw = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    finally:
        workbook.close()

    while raw and raw[-1] is None:
        raw.pop()
    headers = ["" if value is None else str(value).strip() for value in raw]
    if any(not value for value in headers):
        positions = [index + 1 for index, value in enumerate(headers) if not value]
        raise ValueError(f"Blank header cells in {path.name}/{sheet_name}: columns {positions}")
    normalized = [_normalize_name(value) for value in headers]
    duplicates = sorted({value for value in normalized if normalized.count(value) > 1})
    if duplicates:
        raise ValueError(f"Duplicate headers in {path.name}/{sheet_name}: {duplicates}")
    return headers


def _normalize_id(value: Any) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)) and np.isfinite(value) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def _map_binary_label(value: Any) -> int:
    if pd.isna(value):
        raise ValueError("missing label")
    if isinstance(value, (bool, np.bool_)):
        return int(value)
    if isinstance(value, (int, np.integer)) and int(value) in {0, 1}:
        return int(value)
    if isinstance(value, (float, np.floating)) and np.isfinite(value) and float(value) in {0.0, 1.0}:
        return int(value)

    token = _normalize_name(value).replace("_", "-")
    mapping = {
        "responder": 1,
        "response": 1,
        "r": 1,
        "positive": 1,
        "non-responder": 0,
        "nonresponder": 0,
        "nr": 0,
        "negative": 0,
    }
    if token not in mapping:
        raise ValueError(f"unknown label {value!r}")
    return mapping[token]


def _mask_invalid_values(series: pd.Series, rule: FeatureRule) -> pd.Series:
    values = pd.to_numeric(series, errors="coerce").astype(float)
    if rule.negative_is_missing:
        values = values.mask(values < 0, np.nan)
    if rule.zero_is_missing:
        values = values.mask(values == 0, np.nan)
    return values


def load_dataset(spec: DatasetSpec, sheet_name: str) -> DataBundle:
    expected_headers = _read_raw_headers(spec.path, sheet_name)
    frame = pd.read_excel(
        spec.path,
        sheet_name=sheet_name,
        engine="openpyxl",
        na_values=["NA", "N/A", "na", "n/a"],
        keep_default_na=True,
    )
    frame = frame.dropna(how="all").copy()
    frame.columns = [str(column).strip() for column in frame.columns]

    if frame.columns.tolist() != expected_headers:
        raise ValueError(
            f"Header mismatch after reading {spec.path.name}/{sheet_name}. "
            "Check merged cells or duplicate/malformed headers."
        )
    if frame.shape[1] < 3:
        raise ValueError(f"{spec.path.name}/{sheet_name}: expected id, label, and >=1 feature column")
    if frame.columns[0] != ID_COL:
        raise ValueError(
            f"{spec.path.name}/{sheet_name}: first column must be '{ID_COL}', found '{frame.columns[0]}'"
        )

    matched_labels = [column for column in LABEL_COLUMN_CANDIDATES if column in frame.columns]
    if len(matched_labels) != 1:
        raise ValueError(
            f"{spec.path.name}/{sheet_name}: expected exactly one label column from "
            f"{LABEL_COLUMN_CANDIDATES}, found {matched_labels}"
        )
    label_col = matched_labels[0]
    if frame.columns[1] != label_col:
        raise ValueError(
            f"{spec.path.name}/{sheet_name}: label column '{label_col}' must be the second column"
        )

    ids = frame[ID_COL].map(_normalize_id)
    if (ids == "").any():
        rows = (np.flatnonzero((ids == "").to_numpy()) + 2).tolist()[:10]
        raise ValueError(f"Blank IDs in {spec.path.name}/{sheet_name}; Excel rows={rows}")
    if ids.duplicated().any():
        duplicates = ids.loc[ids.duplicated(keep=False)].unique().tolist()[:10]
        raise ValueError(f"Duplicate IDs in {spec.path.name}/{sheet_name}: {duplicates}")

    mapped_labels: List[int] = []
    label_errors: List[str] = []
    for row_number, value in enumerate(frame[label_col].tolist(), start=2):
        try:
            mapped_labels.append(_map_binary_label(value))
        except ValueError as exc:
            label_errors.append(f"row {row_number}: {exc}")
    if label_errors:
        raise ValueError(
            f"Invalid labels in {spec.path.name}/{sheet_name}: " + "; ".join(label_errors[:10])
        )
    y = np.asarray(mapped_labels, dtype=int)
    if set(np.unique(y).tolist()) != {0, 1}:
        raise ValueError(
            f"{spec.path.name}/{sheet_name}: both classes are required; counts={class_counts(y)}"
        )

    feature_names = frame.columns[2:].tolist()
    reserved = {_normalize_name(ID_COL), *(_normalize_name(x) for x in LABEL_COLUMN_CANDIDATES)}
    reserved_hits = [feature for feature in feature_names if _normalize_name(feature) in reserved]
    if reserved_hits:
        raise ValueError(f"Reserved ID/label columns entered the feature block: {reserved_hits}")
    _validate_required_and_forbidden_features(feature_names, spec.dataset_id)

    numeric_columns: Dict[str, pd.Series] = {}
    conversion_errors: List[str] = []
    for feature in feature_names:
        original = frame[feature]
        numeric = pd.to_numeric(original, errors="coerce")
        bad = original.notna() & numeric.isna()
        if bad.any():
            examples = original.loc[bad].astype(str).unique().tolist()[:5]
            conversion_errors.append(f"{feature}: {examples}")
        numeric_columns[feature] = numeric.astype(float)
    if conversion_errors:
        raise ValueError(
            f"Non-numeric feature tokens in {spec.path.name}/{sheet_name}: "
            + "; ".join(conversion_errors[:10])
        )
    X = pd.DataFrame(numeric_columns, index=frame.index)

    rules = resolve_feature_rules(feature_names, spec.dataset_id)
    manifest_rows: List[Dict[str, Any]] = []
    valid_columns: Dict[str, pd.Series] = {}
    for feature in feature_names:
        rule = rules[feature]
        masked = _mask_invalid_values(X[feature], rule)
        valid_columns[feature] = masked
        manifest_rows.append(
            {
                "dataset": spec.key,
                "sheet": sheet_name,
                "feature": feature,
                "transform": rule.transform,
                "zero_is_missing": rule.zero_is_missing,
                "negative_is_missing": rule.negative_is_missing,
                "half_minimum_allowed": rule.half_minimum_allowed,
                "n_original_missing": int(X[feature].isna().sum()),
                "n_zero": int((X[feature] == 0).sum()),
                "n_negative": int((X[feature] < 0).sum()),
                "n_missing_after_rules": int(masked.isna().sum()),
                "missing_rate_after_rules": float(masked.isna().mean()),
                "n_unique_valid": int(masked.nunique(dropna=True)),
            }
        )
    valid_matrix = pd.DataFrame(valid_columns, index=X.index)

    # A feature that is all-missing after the value rules can never be used in
    # any training fold. Excluding only this 100%-missing case globally is
    # leakage-safe and prevents a redundant fail-fast check from pre-empting
    # the training-fold missingness filter below. Features with partial
    # missingness remain in the input universe and are still filtered using
    # training-fold data only.
    all_missing_columns = [
        feature for feature in feature_names if valid_matrix[feature].isna().all()
    ]
    all_missing_set = set(all_missing_columns)
    all_missing_normalized = {
        _normalize_name(feature): feature for feature in all_missing_columns
    }
    required_all_missing = [
        all_missing_normalized[_normalize_name(required)]
        for required in REQUIRED_FEATURES_BY_DATASET.get(spec.dataset_id, ())
        if _normalize_name(required) in all_missing_normalized
    ]
    if required_all_missing:
        raise ValueError(
            f"Required feature(s) are globally all-missing after feature rules in "
            f"{spec.path.name}/{sheet_name}: {required_all_missing}"
        )
    retained_features = [
        feature for feature in feature_names if feature not in all_missing_set
    ]
    if not retained_features:
        raise ValueError(
            f"No usable features remain after feature rules in "
            f"{spec.path.name}/{sheet_name}; all {len(feature_names)} feature(s) "
            "are globally all-missing"
        )

    feature_manifest = pd.DataFrame(manifest_rows)
    feature_manifest["globally_all_missing_after_rules"] = feature_manifest[
        "feature"
    ].isin(all_missing_set)
    feature_manifest["included_in_analysis"] = ~feature_manifest[
        "globally_all_missing_after_rules"
    ]
    feature_manifest["global_exclusion_reason"] = np.where(
        feature_manifest["globally_all_missing_after_rules"],
        "all_missing_after_feature_rules",
        "",
    )

    input_warnings: List[str] = []
    if all_missing_columns:
        shown = all_missing_columns[:20]
        suffix = (
            f"; plus {len(all_missing_columns) - len(shown)} more"
            if len(all_missing_columns) > len(shown)
            else ""
        )
        input_warnings.append(
            f"Excluded {len(all_missing_columns)} globally all-missing feature(s) "
            f"after feature rules: {shown}{suffix}"
        )

    # Keep this row-level QC: a sample with no observed value in any retained
    # feature contains no molecular measurement information at all.
    all_missing_rows = valid_matrix[retained_features].isna().all(axis=1)
    if all_missing_rows.any():
        invalid_ids = ids.loc[all_missing_rows].tolist()[:20]
        raise ValueError(
            f"Samples with all features missing after feature rules in {spec.path.name}/{sheet_name}: "
            f"{invalid_ids}"
        )

    X = X[retained_features].copy()
    rules = {feature: rules[feature] for feature in retained_features}
    X.index = np.arange(len(X))
    ids.index = X.index
    return DataBundle(
        spec=spec,
        sheet_name=sheet_name,
        X=X,
        y=y,
        ids=ids,
        feature_names=retained_features,
        feature_rules=rules,
        feature_manifest=feature_manifest,
        label_col=label_col,
        input_sha256=sha256_file(spec.path),
        input_warnings=input_warnings,
    )


def read_xlsx_dataset(
    path: str,
    sheet_name: str = PRIMARY_SHEET,
    dataset_id: Optional[str] = None,
) -> Tuple[pd.DataFrame, np.ndarray, List[str], pd.Series]:
    """Compatibility loader for downstream scripts.

    New code should use load_dataset() with an explicit DatasetSpec.
    """
    file_path = Path(path).expanduser().resolve()
    match = STANDARD_NAME_RE.match(file_path.name)
    inferred_id = dataset_id
    inferred_timepoint = "pre"
    if match:
        inferred_id = inferred_id or f"dataset{int(match.group('number'))}"
        inferred_timepoint = match.group("timepoint").casefold()
    inferred_id = inferred_id or "dataset1"
    key = f"{inferred_id}_{inferred_timepoint}"
    spec = DatasetSpec(file_path, inferred_id, inferred_timepoint, key, sheet_name)
    bundle = load_dataset(spec, sheet_name)
    return bundle.X, bundle.y, bundle.feature_names, bundle.ids


# =============================================================================
# FOLD-LOCAL PREPROCESSING
# =============================================================================

class MetaboPreprocessor:
    """Training-fold-only preprocessing with feature-specific value rules.

    half_minimum:
        mask invalid values -> training half-minimum (or median for signed
        identity features) -> feature transform -> training z-score

    knn:
        mask invalid values -> feature transform -> training NaN-aware z-score
        -> KNNImputer fitted on the standardized training fold
    """

    def __init__(
        self,
        feature_names: Sequence[str],
        missing_rate_threshold: float = MISSING_RATE_THRESHOLD,
        feature_rules: Optional[Mapping[str, FeatureRule]] = None,
        imputation: str = "half_minimum",
        knn_neighbors: int = KNN_NEIGHBORS,
    ):
        self.feature_names = list(feature_names)
        self.missing_rate_threshold = float(missing_rate_threshold)
        self.feature_rules = {
            feature: (feature_rules or {}).get(feature, _default_feature_rule(feature))
            for feature in self.feature_names
        }
        self.imputation = str(imputation)
        self.knn_neighbors = int(knn_neighbors)

        if self.imputation not in {"half_minimum", "knn"}:
            raise ValueError(f"Unsupported imputation strategy: {self.imputation}")
        if not (0.0 <= self.missing_rate_threshold < 1.0):
            raise ValueError("missing_rate_threshold must be in [0, 1)")

        self._kept_after_missing: List[str] = []
        self._kept_features: List[str] = []
        self._missing_rate: Optional[pd.Series] = None
        self._impute_values_raw: Optional[pd.Series] = None
        self._mean_: Optional[pd.Series] = None
        self._std_: Optional[pd.Series] = None
        self._knn_imputer: Optional[KNNImputer] = None
        self._knn_neighbors_used: Optional[int] = None

    def _numeric_frame(self, X: pd.DataFrame) -> pd.DataFrame:
        missing_columns = [feature for feature in self.feature_names if feature not in X.columns]
        if missing_columns:
            raise ValueError(f"Input is missing required feature columns: {missing_columns[:20]}")
        return pd.DataFrame(
            {
                feature: pd.to_numeric(X[feature], errors="coerce").astype(float)
                for feature in self.feature_names
            },
            index=X.index,
        )

    def _mask_raw(self, X: pd.DataFrame) -> pd.DataFrame:
        numeric = self._numeric_frame(X)
        return pd.DataFrame(
            {
                feature: _mask_invalid_values(numeric[feature], self.feature_rules[feature])
                for feature in self.feature_names
            },
            index=numeric.index,
        )

    def _transform_observed(self, X: pd.DataFrame) -> pd.DataFrame:
        columns: Dict[str, pd.Series] = {}
        for feature in X.columns:
            rule = self.feature_rules[feature]
            values = X[feature].astype(float)
            if rule.transform == "log1p":
                invalid = values.notna() & (values < 0)
                if invalid.any():
                    raise ValueError(f"{feature}: negative values reached log1p after masking")
                columns[feature] = np.log1p(values)
            elif rule.transform == "identity":
                columns[feature] = values
            else:  # defensive; validated when rules are built
                raise ValueError(f"{feature}: unsupported transform {rule.transform}")
        return pd.DataFrame(columns, index=X.index)

    def fit(self, X: pd.DataFrame) -> "MetaboPreprocessor":
        raw = self._mask_raw(X)
        missing_rate = raw.isna().mean(axis=0)
        self._missing_rate = missing_rate

        kept = [
            feature
            for feature in self.feature_names
            if missing_rate[feature] <= self.missing_rate_threshold and not raw[feature].isna().all()
        ]
        if not kept:
            raise ValueError(
                "No features remain after the training-fold missingness filter; "
                f"threshold={self.missing_rate_threshold}"
            )
        self._kept_after_missing = kept
        raw_kept = raw[kept].copy()

        if self.imputation == "half_minimum":
            impute_values: Dict[str, float] = {}
            for feature in kept:
                observed = raw_kept[feature].dropna().astype(float)
                if observed.empty:
                    raise ValueError(f"{feature}: no observed values in training fold")
                rule = self.feature_rules[feature]
                if rule.half_minimum_allowed:
                    positive = observed[observed > 0]
                    if positive.empty:
                        raise ValueError(
                            f"{feature}: half-minimum imputation requires a positive training value"
                        )
                    impute_values[feature] = float(positive.min() / 2.0)
                else:
                    impute_values[feature] = float(observed.median())

            self._impute_values_raw = pd.Series(impute_values, dtype=float)
            filled = raw_kept.fillna(self._impute_values_raw)
            transformed = self._transform_observed(filled)
            variance = transformed.var(axis=0, ddof=0)
            used = [feature for feature in kept if np.isfinite(variance[feature]) and variance[feature] > 0]
            if not used:
                raise ValueError("All training-fold features are constant after preprocessing")
            self._kept_features = used
            used_frame = transformed[used]
            self._mean_ = used_frame.mean(axis=0)
            self._std_ = used_frame.std(axis=0, ddof=0)
            if (self._std_ <= 0).any() or self._std_.isna().any():
                raise ValueError("Invalid training-fold standard deviations after constant filtering")
        else:
            transformed = self._transform_observed(raw_kept)
            mean = transformed.mean(axis=0, skipna=True)
            std = transformed.std(axis=0, ddof=0, skipna=True)
            used = [
                feature
                for feature in kept
                if np.isfinite(mean[feature]) and np.isfinite(std[feature]) and std[feature] > 0
            ]
            if not used:
                raise ValueError("All training-fold features are constant before KNN imputation")
            self._kept_features = used
            self._mean_ = mean[used]
            self._std_ = std[used]
            standardized = (transformed[used] - self._mean_) / self._std_
            neighbors = min(self.knn_neighbors, len(standardized))
            if neighbors < 1:
                raise ValueError("KNN imputation received an empty training fold")
            self._knn_neighbors_used = int(neighbors)
            self._knn_imputer = KNNImputer(n_neighbors=neighbors, weights="uniform")
            fitted = self._knn_imputer.fit_transform(standardized)
            if fitted.shape[1] != len(used) or not np.isfinite(fitted).all():
                raise RuntimeError("KNN imputation produced an invalid training matrix")

        return self

    def transform(self, X: pd.DataFrame) -> np.ndarray:
        if self._mean_ is None or self._std_ is None or not self._kept_features:
            raise RuntimeError("MetaboPreprocessor is not fitted")
        raw = self._mask_raw(X)[self._kept_after_missing]

        if self.imputation == "half_minimum":
            if self._impute_values_raw is None:
                raise RuntimeError("Missing fitted half-minimum imputation values")
            filled = raw.fillna(self._impute_values_raw)
            transformed = self._transform_observed(filled)[self._kept_features]
            output = (transformed - self._mean_) / self._std_
            array = output.to_numpy(dtype=float)
        else:
            if self._knn_imputer is None:
                raise RuntimeError("Missing fitted KNN imputer")
            transformed = self._transform_observed(raw)[self._kept_features]
            standardized = (transformed - self._mean_) / self._std_
            array = self._knn_imputer.transform(standardized)

        if array.shape[1] != len(self._kept_features) or not np.isfinite(array).all():
            raise RuntimeError("Preprocessing produced non-finite values or unexpected columns")
        return np.asarray(array, dtype=float)

    @property
    def kept_features_(self) -> List[str]:
        return list(self._kept_features)

    @property
    def kept_after_missing_(self) -> List[str]:
        return list(self._kept_after_missing)

    def audit_table(self) -> pd.DataFrame:
        rows: List[Dict[str, Any]] = []
        for feature in self.feature_names:
            rule = self.feature_rules[feature]
            rows.append(
                {
                    "feature": feature,
                    "transform": rule.transform,
                    "zero_is_missing": rule.zero_is_missing,
                    "negative_is_missing": rule.negative_is_missing,
                    "imputation": self.imputation,
                    "training_missing_rate": (
                        float(self._missing_rate[feature])
                        if self._missing_rate is not None and feature in self._missing_rate
                        else np.nan
                    ),
                    "kept_after_missing_filter": feature in self._kept_after_missing,
                    "model_used": feature in self._kept_features,
                    "impute_value_raw": (
                        float(self._impute_values_raw[feature])
                        if self._impute_values_raw is not None and feature in self._impute_values_raw
                        else np.nan
                    ),
                    "training_mean_transformed": (
                        float(self._mean_[feature])
                        if self._mean_ is not None and feature in self._mean_
                        else np.nan
                    ),
                    "training_sd_transformed": (
                        float(self._std_[feature])
                        if self._std_ is not None and feature in self._std_
                        else np.nan
                    ),
                    "knn_neighbors_used": self._knn_neighbors_used,
                }
            )
        return pd.DataFrame(rows)


# =============================================================================
# MODEL FITTING AND INNER-LOOP SELECTION
# =============================================================================

def _sklearn_version_tuple() -> Tuple[int, int]:
    match = re.match(r"^(\d+)\.(\d+)", sklearn.__version__)
    return (int(match.group(1)), int(match.group(2))) if match else (0, 0)


def _make_logistic_model(
    kind: str,
    C: float,
    l1_ratio: Optional[float],
    max_iter: int,
) -> LogisticRegression:
    common: Dict[str, Any] = {
        "C": float(C),
        "class_weight": "balanced",
        "max_iter": int(max_iter),
        "tol": float(TOL),
        "random_state": int(RANDOM_STATE),
    }
    modern = _sklearn_version_tuple() >= (1, 8)
    if kind == "elasticnet":
        if l1_ratio is None:
            raise ValueError("elasticnet requires l1_ratio")
        common.update({"solver": "saga", "l1_ratio": float(l1_ratio)})
        if not modern:
            common["penalty"] = "elasticnet"
    elif kind == "l2":
        if modern:
            common.update({"solver": "lbfgs", "l1_ratio": 0.0})
        else:
            common.update({"solver": "liblinear", "penalty": "l2"})
    else:
        raise ValueError(f"Unknown logistic-regression kind: {kind}")
    return LogisticRegression(**common)


def fit_logreg_with_refit(
    X: np.ndarray,
    y: np.ndarray,
    kind: str,
    C: float,
    l1_ratio: Optional[float],
    dataset_name: str,
    stage: str,
    log_lines: List[str],
    warning_rows: Optional[List[Dict[str, Any]]] = None,
) -> LogisticRegression:
    if kind == "elasticnet":
        max_iter = EN_MAX_ITER_INIT
        limit = EN_MAX_ITER_LIMIT
    elif kind == "l2":
        max_iter = L2_MAX_ITER_INIT
        limit = L2_MAX_ITER_LIMIT
    else:
        raise ValueError(f"Unknown model kind: {kind}")

    while True:
        model = _make_logistic_model(kind, C, l1_ratio, max_iter)
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            model.fit(X, y)

        convergence = [item for item in caught if issubclass(item.category, ConvergenceWarning)]
        other_warnings = [item for item in caught if not issubclass(item.category, ConvergenceWarning)]
        for item in other_warnings:
            message = f"[{item.category.__name__}] {dataset_name} | {stage}: {item.message}"
            log_lines.append(message)
            if warning_rows is not None:
                warning_rows.append(
                    {
                        "dataset": dataset_name,
                        "stage": stage,
                        "warning_type": item.category.__name__,
                        "message": str(item.message),
                    }
                )
            warnings.warn(message, item.category, stacklevel=2)

        if not convergence:
            return model
        next_iter = min(max_iter * 2, limit)
        message = (
            f"[WARN] {dataset_name} | {stage}: convergence warning at max_iter={max_iter}; "
            f"next_max_iter={next_iter}"
        )
        log_lines.append(message)
        if warning_rows is not None:
            warning_rows.append(
                {
                    "dataset": dataset_name,
                    "stage": stage,
                    "warning_type": "ConvergenceWarning",
                    "message": message,
                }
            )
        if max_iter >= limit:
            if HARD_FAIL_ON_NONCONVERGENCE:
                raise RuntimeError(
                    f"{dataset_name} | {stage}: logistic regression did not converge by max_iter={limit}"
                )
            return model
        max_iter = next_iter


def _coerce_cv_plan(
    cv: Any,
    X: pd.DataFrame,
    y: np.ndarray,
    expected_repeats: Optional[int] = None,
) -> CVPlan:
    if isinstance(cv, CVPlan):
        return cv
    splits = list(cv.split(X, y))
    if not splits:
        raise ValueError("CV object produced no splits")
    total = len(splits)
    repeats = int(expected_repeats or 1)
    if total % repeats != 0:
        repeats = 1
    n_splits = total // repeats
    records = [
        SplitRecord(
            repeat=index // n_splits + 1,
            fold=index % n_splits + 1,
            train_idx=np.asarray(train, dtype=int),
            test_idx=np.asarray(test, dtype=int),
        )
        for index, (train, test) in enumerate(splits)
    ]
    return CVPlan(records, n_splits=n_splits, n_repeats=repeats)


def _make_preprocessor(
    features: Sequence[str],
    feature_rules: Mapping[str, FeatureRule],
    imputation: str,
) -> MetaboPreprocessor:
    return MetaboPreprocessor(
        feature_names=features,
        missing_rate_threshold=MISSING_RATE_THRESHOLD,
        feature_rules={feature: feature_rules[feature] for feature in features},
        imputation=imputation,
        knn_neighbors=KNN_NEIGHBORS,
    )


def tune_elasticnet(
    X: pd.DataFrame,
    y: np.ndarray,
    feature_names: List[str],
    inner_cv: Any,
    dataset_name: str,
    outer_iter: int,
    log_lines: List[str],
    feature_rules: Optional[Mapping[str, FeatureRule]] = None,
    imputation: str = "half_minimum",
    warning_rows: Optional[List[Dict[str, Any]]] = None,
) -> Tuple[float, float, pd.DataFrame]:
    rules = dict(feature_rules or resolve_feature_rules(feature_names, None))
    plan = _coerce_cv_plan(inner_cv, X, y, expected_repeats=INNER_REPEATS)
    rows: List[Dict[str, Any]] = []
    best_mean = -np.inf
    best_params = (float(EN_C_GRID[0]), float(EN_L1_RATIO_GRID[0]))

    for C in EN_C_GRID:
        for l1_ratio in EN_L1_RATIO_GRID:
            fold_rows: List[Dict[str, Any]] = []
            for record in plan.records:
                X_train = X.iloc[record.train_idx]
                X_valid = X.iloc[record.test_idx]
                y_train = y[record.train_idx]
                y_valid = y[record.test_idx]
                preprocessor = _make_preprocessor(feature_names, rules, imputation).fit(X_train)
                train_matrix = preprocessor.transform(X_train)
                valid_matrix = preprocessor.transform(X_valid)
                model = fit_logreg_with_refit(
                    train_matrix,
                    y_train,
                    kind="elasticnet",
                    C=float(C),
                    l1_ratio=float(l1_ratio),
                    dataset_name=dataset_name,
                    stage=(
                        f"EN-tune outer={outer_iter} inner_repeat={record.repeat} "
                        f"inner_fold={record.fold}"
                    ),
                    log_lines=log_lines,
                    warning_rows=warning_rows,
                )
                probability = model.predict_proba(valid_matrix)[:, 1]
                fold_rows.append(
                    {
                        "repeat": record.repeat,
                        "fold": record.fold,
                        "auc": float(roc_auc_score(y_valid, probability)),
                    }
                )

            fold_df = pd.DataFrame(fold_rows)
            repeat_means = fold_df.groupby("repeat", sort=True)["auc"].mean()
            mean_auc = float(repeat_means.mean())
            rows.append(
                {
                    "C": float(C),
                    "l1_ratio": float(l1_ratio),
                    "mean_auc": mean_auc,
                    "sd_repeat_mean_auc": (
                        float(repeat_means.std(ddof=1)) if len(repeat_means) > 1 else np.nan
                    ),
                    "n_inner_repeats": int(len(repeat_means)),
                    "n_inner_evaluations": int(len(fold_df)),
                }
            )
            if mean_auc > best_mean:
                best_mean = mean_auc
                best_params = (float(C), float(l1_ratio))

    result = pd.DataFrame(rows).sort_values(
        ["mean_auc", "C", "l1_ratio"], ascending=[False, True, True]
    ).reset_index(drop=True)
    message = (
        f"[INFO] {dataset_name} | outer={outer_iter}: best EN C={best_params[0]}, "
        f"l1_ratio={best_params[1]}, inner mean AUC={best_mean:.4f}"
    )
    print(message)
    log_lines.append(message)
    return best_params[0], best_params[1], result


def stability_select_and_rank(
    X: pd.DataFrame,
    y: np.ndarray,
    feature_names: List[str],
    inner_cv: Any,
    en_C: float,
    en_l1: float,
    dataset_name: str,
    outer_iter: int,
    log_lines: List[str],
    feature_rules: Optional[Mapping[str, FeatureRule]] = None,
    imputation: str = "half_minimum",
    warning_rows: Optional[List[Dict[str, Any]]] = None,
) -> Tuple[List[str], pd.DataFrame]:
    rules = dict(feature_rules or resolve_feature_rules(feature_names, None))
    plan = _coerce_cv_plan(inner_cv, X, y, expected_repeats=INNER_REPEATS)
    position = {feature: index for index, feature in enumerate(feature_names)}
    p = len(feature_names)
    eligible_count = np.zeros(p, dtype=int)
    nonzero_count = np.zeros(p, dtype=int)
    abs_coefficient_sum = np.zeros(p, dtype=float)
    signed_coefficient_sum = np.zeros(p, dtype=float)

    for record in plan.records:
        X_train = X.iloc[record.train_idx]
        y_train = y[record.train_idx]
        preprocessor = _make_preprocessor(feature_names, rules, imputation).fit(X_train)
        train_matrix = preprocessor.transform(X_train)
        model = fit_logreg_with_refit(
            train_matrix,
            y_train,
            kind="elasticnet",
            C=en_C,
            l1_ratio=en_l1,
            dataset_name=dataset_name,
            stage=(
                f"EN-stability outer={outer_iter} inner_repeat={record.repeat} "
                f"inner_fold={record.fold}"
            ),
            log_lines=log_lines,
            warning_rows=warning_rows,
        )
        coefficients = model.coef_.ravel()
        for feature, coefficient in zip(preprocessor.kept_features_, coefficients):
            index = position[feature]
            eligible_count[index] += 1
            nonzero_count[index] += int(coefficient != 0.0)
            abs_coefficient_sum[index] += abs(float(coefficient))
            signed_coefficient_sum[index] += float(coefficient)

    total = len(plan.records)
    table = pd.DataFrame(
        {
            "feature": feature_names,
            "eligible_freq": eligible_count / float(total),
            "nonzero_freq": nonzero_count / float(total),
            "mean_abs_coef": abs_coefficient_sum / float(total),
            "mean_signed_coef": signed_coefficient_sum / float(total),
        }
    )
    table = table.loc[table["eligible_freq"] > 0].sort_values(
        ["nonzero_freq", "mean_abs_coef", "feature"],
        ascending=[False, False, True],
    ).reset_index(drop=True)
    table.insert(0, "rank", np.arange(1, len(table) + 1, dtype=int))
    ranking = table["feature"].tolist()
    if not ranking:
        raise ValueError(f"{dataset_name}: no feature was eligible in inner resampling")
    return ranking, table


def select_k_by_1se(
    X: pd.DataFrame,
    y: np.ndarray,
    ranking: List[str],
    inner_cv: Any,
    dataset_name: str,
    outer_iter: int,
    log_lines: List[str],
    feature_rules: Optional[Mapping[str, FeatureRule]] = None,
    imputation: str = "half_minimum",
    max_k: Optional[int] = None,
    warning_rows: Optional[List[Dict[str, Any]]] = None,
) -> Tuple[int, float, pd.DataFrame]:
    rules = dict(feature_rules or resolve_feature_rules(ranking, None))
    plan = _coerce_cv_plan(inner_cv, X, y, expected_repeats=INNER_REPEATS)
    limit = min(int(MAX_K if max_k is None else max_k), len(ranking))
    if limit < 1:
        raise ValueError("No ranked feature is available for k selection")

    rows: List[Dict[str, Any]] = []
    for k in range(1, limit + 1):
        features = ranking[:k]
        best_mean = -np.inf
        best_C = float(L2_C_GRID[0])
        best_repeat_means = pd.Series(dtype=float)

        for C in L2_C_GRID:
            fold_rows: List[Dict[str, Any]] = []
            for record in plan.records:
                X_train = X.iloc[record.train_idx][features]
                X_valid = X.iloc[record.test_idx][features]
                y_train = y[record.train_idx]
                y_valid = y[record.test_idx]
                preprocessor = _make_preprocessor(features, rules, imputation).fit(X_train)
                train_matrix = preprocessor.transform(X_train)
                valid_matrix = preprocessor.transform(X_valid)
                model = fit_logreg_with_refit(
                    train_matrix,
                    y_train,
                    kind="l2",
                    C=float(C),
                    l1_ratio=None,
                    dataset_name=dataset_name,
                    stage=(
                        f"L2-kselect outer={outer_iter} k={k} C={C} "
                        f"inner_repeat={record.repeat} inner_fold={record.fold}"
                    ),
                    log_lines=log_lines,
                    warning_rows=warning_rows,
                )
                probability = model.predict_proba(valid_matrix)[:, 1]
                fold_rows.append(
                    {
                        "repeat": record.repeat,
                        "fold": record.fold,
                        "auc": float(roc_auc_score(y_valid, probability)),
                    }
                )

            fold_df = pd.DataFrame(fold_rows)
            repeat_means = fold_df.groupby("repeat", sort=True)["auc"].mean()
            mean_auc = float(repeat_means.mean())
            if mean_auc > best_mean:
                best_mean = mean_auc
                best_C = float(C)
                best_repeat_means = repeat_means.copy()

        repeat_sd = (
            float(best_repeat_means.std(ddof=1)) if len(best_repeat_means) > 1 else 0.0
        )
        repeat_se = repeat_sd / math.sqrt(len(best_repeat_means)) if len(best_repeat_means) else np.inf
        rows.append(
            {
                "outer_iter": int(outer_iter),
                "k": int(k),
                "best_C": best_C,
                "mean_auc": best_mean,
                "sd_repeat_mean_auc": repeat_sd,
                "se_auc": repeat_se,
                "se_unit": "inner_repeat_mean",
                "n_inner_repeats": int(len(best_repeat_means)),
                "n_inner_evaluations": int(len(plan.records)),
            }
        )

    table = pd.DataFrame(rows)
    best_index = table["mean_auc"].idxmax()
    best_mean = float(table.loc[best_index, "mean_auc"])
    best_se = float(table.loc[best_index, "se_auc"])
    threshold = best_mean - best_se
    eligible = table.loc[table["mean_auc"] >= threshold].sort_values("k")
    chosen = eligible.iloc[0]
    chosen_k = int(chosen["k"])
    chosen_C = float(chosen["best_C"])
    message = (
        f"[INFO] {dataset_name} | outer={outer_iter}: 1-SE threshold={threshold:.4f}; "
        f"chosen k={chosen_k}, L2 C={chosen_C}"
    )
    print(message)
    log_lines.append(message)
    return chosen_k, chosen_C, table


# =============================================================================
# METRICS AND FEATURE-STABILITY SUMMARIES
# =============================================================================

def bootstrap_auc_ci(
    y_true: np.ndarray,
    probability: np.ndarray,
    n_boot: int,
    seed: int = RANDOM_STATE,
) -> Tuple[float, float]:
    """Stratified patient bootstrap; descriptive, not CV-variance corrected."""
    y = np.asarray(y_true, dtype=int)
    p = np.asarray(probability, dtype=float)
    positive = np.flatnonzero(y == 1)
    negative = np.flatnonzero(y == 0)
    if not len(positive) or not len(negative):
        raise ValueError("Both classes are required for an AUROC bootstrap")
    rng = np.random.default_rng(seed)
    values = np.empty(int(n_boot), dtype=float)
    for index in range(int(n_boot)):
        sample = np.concatenate(
            [
                rng.choice(positive, size=len(positive), replace=True),
                rng.choice(negative, size=len(negative), replace=True),
            ]
        )
        values[index] = roc_auc_score(y[sample], p[sample])
    low, high = np.percentile(values, [2.5, 97.5])
    return float(low), float(high)


def nogueira_stability(
    selected_sets: Sequence[Sequence[str]],
    feature_universe: Sequence[str],
) -> Dict[str, Any]:
    features = list(dict.fromkeys(feature_universe))
    index = {feature: position for position, feature in enumerate(features)}
    B = len(selected_sets)
    p = len(features)
    if B < 2 or p == 0:
        return {"nogueira_stability": np.nan, "stability_status": "insufficient_B_or_p"}

    matrix = np.zeros((B, p), dtype=np.uint8)
    for row, selected in enumerate(selected_sets):
        selected_unique = set(selected)
        unknown = selected_unique.difference(index)
        if unknown:
            raise ValueError(f"Unknown selected features in stability calculation: {sorted(unknown)}")
        for feature in selected_unique:
            matrix[row, index[feature]] = 1

    k_values = matrix.sum(axis=1)
    mean_k = float(k_values.mean())
    q = mean_k / float(p)
    base = {
        "n_outer_fits": int(B),
        "feature_universe_n": int(p),
        "mean_model_feature_n": mean_k,
        "median_model_feature_n": float(np.median(k_values)),
        "min_model_feature_n": int(k_values.min()),
        "max_model_feature_n": int(k_values.max()),
    }
    if np.isclose(q, 0.0) or np.isclose(q, 1.0):
        return {
            **base,
            "nogueira_stability": np.nan,
            "stability_status": "degenerate_all_none_or_all",
        }

    frequency = matrix.mean(axis=0)
    sample_variance = B / float(B - 1) * frequency * (1.0 - frequency)
    value = 1.0 - float(sample_variance.mean()) / (q * (1.0 - q))
    return {**base, "nogueira_stability": value, "stability_status": "ok"}


def pairwise_jaccard_table(
    selected_records: Sequence[Tuple[int, int, Sequence[str]]],
) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    for left, right in combinations(selected_records, 2):
        repeat_a, fold_a, features_a = left
        repeat_b, fold_b, features_b = right
        set_a = set(features_a)
        set_b = set(features_b)
        union = set_a | set_b
        value = 1.0 if not union else len(set_a & set_b) / float(len(union))
        rows.append(
            {
                "repeat_a": repeat_a,
                "fold_a": fold_a,
                "repeat_b": repeat_b,
                "fold_b": fold_b,
                "same_repeat": repeat_a == repeat_b,
                "jaccard": float(value),
            }
        )
    return pd.DataFrame(rows)


def summarize_jaccard(table: pd.DataFrame) -> Dict[str, Any]:
    if table.empty:
        return {
            "jaccard_n_pairs": 0,
            "jaccard_median": np.nan,
            "jaccard_q1": np.nan,
            "jaccard_q3": np.nan,
            "jaccard_mean": np.nan,
            "jaccard_sd": np.nan,
            "jaccard_min": np.nan,
            "jaccard_max": np.nan,
        }
    values = table["jaccard"].to_numpy(dtype=float)
    return {
        "jaccard_n_pairs": int(len(values)),
        "jaccard_median": float(np.median(values)),
        "jaccard_q1": float(np.quantile(values, 0.25)),
        "jaccard_q3": float(np.quantile(values, 0.75)),
        "jaccard_mean": float(np.mean(values)),
        "jaccard_sd": float(np.std(values, ddof=1)) if len(values) > 1 else np.nan,
        "jaccard_min": float(np.min(values)),
        "jaccard_max": float(np.max(values)),
    }


def _fold_count_rows(
    dataset_key: str,
    scenario_name: str,
    level: str,
    outer_record: SplitRecord,
    y_outer_train: np.ndarray,
    inner_plan: Optional[CVPlan] = None,
) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    if level == "outer":
        raise ValueError("Use _outer_fold_count_rows for outer counts")
    if inner_plan is None:
        return rows
    for inner_record in inner_plan.records:
        for subset, indices in (
            ("train", inner_record.train_idx),
            ("validation", inner_record.test_idx),
        ):
            counts = class_counts(y_outer_train[indices])
            rows.append(
                {
                    "dataset": dataset_key,
                    "scenario": scenario_name,
                    "level": "inner",
                    "outer_repeat": outer_record.repeat,
                    "outer_fold": outer_record.fold,
                    "inner_repeat": inner_record.repeat,
                    "inner_fold": inner_record.fold,
                    "subset": subset,
                    **counts,
                }
            )
    return rows


def _outer_fold_count_rows(
    dataset_key: str,
    scenario_name: str,
    record: SplitRecord,
    y: np.ndarray,
) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for subset, indices in (("train", record.train_idx), ("test", record.test_idx)):
        counts = class_counts(y[indices])
        rows.append(
            {
                "dataset": dataset_key,
                "scenario": scenario_name,
                "level": "outer",
                "outer_repeat": record.repeat,
                "outer_fold": record.fold,
                "inner_repeat": np.nan,
                "inner_fold": np.nan,
                "subset": subset,
                **counts,
            }
        )
    return rows


def _output_dir_for(spec: DatasetSpec, scenario: ScenarioSpec) -> Path:
    root = OUTPUT_ROOT / RUN_TAG if RUN_TAG else OUTPUT_ROOT
    if scenario.name == "primary" and spec.dataset_id != "dataset6":
        section = "main"
    elif scenario.name == "primary" and spec.dataset_id == "dataset6":
        section = "supplement/dataset6_small_sample"
    else:
        section = "supplement/sensitivity"
    return root / section / scenario.name / spec.key


def _selector_provenance() -> Dict[str, str]:
    selector_path = Path(dataset_selection.__file__).resolve()
    return {
        "dataset_selection_path": str(selector_path),
        "dataset_selection_sha256": sha256_file(selector_path),
    }


def _scenario_settings_payload(spec: DatasetSpec, scenario: ScenarioSpec, outer_splits_used: int) -> Dict[str, Any]:
    return {
        **_selector_provenance(),
        "dataset": asdict(spec),
        "scenario": asdict(scenario),
        "invocation_scenario": ACTIVE_SCENARIO,
        "invocation_scenario_metadata": ACTIVE_SCENARIO_METADATA,
        "model_sequence": [
            "elastic_net_tuning",
            "inner_resampling_feature_ranking",
            "top_k_and_l2_C_1SE_selection",
            "outer_test_prediction",
        ],
        "random_state": RANDOM_STATE,
        "outer_splits_used": outer_splits_used,
        "outer_repeats": OUTER_REPEATS,
        "inner_splits_target": INNER_SPLITS_TARGET,
        "inner_repeats": INNER_REPEATS,
        "missing_rate_threshold": MISSING_RATE_THRESHOLD,
        "global_all_missing_feature_policy": (
            "exclude before CV and retain in feature_manifest; required features still fail"
        ),
        "partial_missing_feature_policy": "filter within each training fold only",
        "knn_neighbors": KNN_NEIGHBORS,
        "class_weight": "balanced",
        "en_C_grid": EN_C_GRID,
        "en_l1_ratio_grid": EN_L1_RATIO_GRID,
        "l2_C_grid": L2_C_GRID,
        "k_1se_standard_error_unit": "mean AUC for each inner repeat",
        "threshold_dependent_metrics": "not computed; no prespecified clinical threshold",
        "inner_k_curve_note": (
            "Feature ranking and k evaluation reuse the outer-training inner resamples. "
            "The inner k curve is a tuning diagnostic, not an unbiased performance estimate; "
            "formal discrimination comes only from outer held-out predictions."
        ),
    }


def run_nested_cv(spec: DatasetSpec, scenario: ScenarioSpec) -> RunResult:
    scenario.validate()
    started = time.time()
    started_iso = datetime.now(timezone.utc).isoformat()
    output_dir = _output_dir_for(spec, scenario)
    prepare_output_dir(output_dir)

    log_lines: List[str] = []
    warning_rows: List[Dict[str, Any]] = []
    label = f"{spec.key} | {scenario.name}"
    log_lines.extend(
        [
            f"[INFO] dataset={spec.key}",
            f"[INFO] scenario={scenario.name}",
            f"[INFO] input={spec.path}",
            f"[INFO] sheet={scenario.sheet_name}",
        ]
    )
    print("=" * 88)
    print(f"[START] {label}")
    print("=" * 88)

    bundle = load_dataset(spec, scenario.sheet_name)
    X, y, ids = bundle.X, bundle.y, bundle.ids
    feature_names = bundle.feature_names
    rules = bundle.feature_rules
    n, p = X.shape
    counts = class_counts(y)
    input_feature_count = int(len(bundle.feature_manifest))
    excluded_global_count = int(
        bundle.feature_manifest["globally_all_missing_after_rules"].sum()
    )
    for message in bundle.input_warnings:
        formatted = f"[WARN] {label} | input-feature-filter: {message}"
        print(formatted)
        log_lines.append(formatted)
        warning_rows.append(
            {
                "dataset": label,
                "stage": "input-feature-filter",
                "warning_type": "GloballyAllMissingFeatureExcluded",
                "message": message,
            }
        )
    log_lines.append(
        f"[INFO] samples={n}, input_features={input_feature_count}, "
        f"analysis_features={p}, globally_excluded_features={excluded_global_count}, "
        f"class_counts={counts}"
    )

    outer_plan = make_cv_plan(
        y,
        target_splits=scenario.outer_splits_target,
        n_repeats=OUTER_REPEATS,
        context=f"{label} outer",
        log_lines=log_lines,
        random_state=RANDOM_STATE,
    )

    prediction_rows: List[Dict[str, Any]] = []
    parameter_rows: List[Dict[str, Any]] = []
    k_curve_frames: List[pd.DataFrame] = []
    fold_count_rows: List[Dict[str, Any]] = []
    outer_feature_rows: List[Dict[str, Any]] = []
    selected_records: List[Tuple[int, int, Sequence[str]]] = []

    for outer_iter, outer_record in enumerate(outer_plan.records, start=1):
        train_idx = outer_record.train_idx
        test_idx = outer_record.test_idx
        X_train = X.iloc[train_idx]
        X_test = X.iloc[test_idx]
        y_train = y[train_idx]
        y_test = y[test_idx]

        fold_count_rows.extend(_outer_fold_count_rows(spec.key, scenario.name, outer_record, y))
        inner_plan = make_cv_plan(
            y_train,
            target_splits=INNER_SPLITS_TARGET,
            n_repeats=INNER_REPEATS,
            context=(
                f"{label} inner outer_repeat={outer_record.repeat} "
                f"outer_fold={outer_record.fold}"
            ),
            log_lines=log_lines,
            random_state=RANDOM_STATE,
        )
        fold_count_rows.extend(
            _fold_count_rows(
                spec.key,
                scenario.name,
                "inner",
                outer_record,
                y_train,
                inner_plan,
            )
        )

        en_C, en_l1, _ = tune_elasticnet(
            X_train,
            y_train,
            feature_names,
            inner_plan,
            dataset_name=label,
            outer_iter=outer_iter,
            log_lines=log_lines,
            feature_rules=rules,
            imputation=scenario.imputation,
            warning_rows=warning_rows,
        )
        ranking, stability_table = stability_select_and_rank(
            X_train,
            y_train,
            feature_names,
            inner_plan,
            en_C=en_C,
            en_l1=en_l1,
            dataset_name=label,
            outer_iter=outer_iter,
            log_lines=log_lines,
            feature_rules=rules,
            imputation=scenario.imputation,
            warning_rows=warning_rows,
        )

        # Limit the ranked universe to features usable when fitting this entire
        # outer-training set. This prevents alphabetic/tie fallback features
        # that are not actually estimable from entering Top-k.
        eligibility_preprocessor = _make_preprocessor(
            feature_names, rules, scenario.imputation
        ).fit(X_train)
        outer_eligible = set(eligibility_preprocessor.kept_features_)
        ranking = [feature for feature in ranking if feature in outer_eligible]
        stability_table = stability_table.loc[
            stability_table["feature"].isin(outer_eligible)
        ].copy()
        stability_table = stability_table.sort_values("rank").reset_index(drop=True)
        stability_table["rank"] = np.arange(1, len(stability_table) + 1, dtype=int)
        if not ranking:
            raise ValueError(f"{label}: no outer-training eligible ranked feature")

        k_selected, l2_C, k_table = select_k_by_1se(
            X_train,
            y_train,
            ranking,
            inner_plan,
            dataset_name=label,
            outer_iter=outer_iter,
            log_lines=log_lines,
            feature_rules=rules,
            imputation=scenario.imputation,
            max_k=scenario.max_k,
            warning_rows=warning_rows,
        )
        k_table.insert(0, "dataset", spec.key)
        k_table.insert(1, "scenario", scenario.name)
        k_table["outer_repeat"] = outer_record.repeat
        k_table["outer_fold"] = outer_record.fold
        k_curve_frames.append(k_table)

        selected = ranking[:k_selected]
        preprocessor = _make_preprocessor(selected, rules, scenario.imputation).fit(
            X_train[selected]
        )
        train_matrix = preprocessor.transform(X_train[selected])
        test_matrix = preprocessor.transform(X_test[selected])
        model = fit_logreg_with_refit(
            train_matrix,
            y_train,
            kind="l2",
            C=l2_C,
            l1_ratio=None,
            dataset_name=label,
            stage=f"L2-outerfit repeat={outer_record.repeat} fold={outer_record.fold}",
            log_lines=log_lines,
            warning_rows=warning_rows,
        )
        probability = model.predict_proba(test_matrix)[:, 1]
        if not np.isfinite(probability).all():
            raise RuntimeError(f"{label}: non-finite outer-test probabilities")

        for local_index, original_index in enumerate(test_idx):
            prediction_rows.append(
                {
                    "dataset": spec.key,
                    "scenario": scenario.name,
                    ID_COL: ids.iloc[original_index],
                    "repeat": outer_record.repeat,
                    "outer_fold": outer_record.fold,
                    "y_true": int(y[original_index]),
                    "p": float(probability[local_index]),
                }
            )

        model_features = preprocessor.kept_features_
        coefficient_map = {
            feature: float(coefficient)
            for feature, coefficient in zip(model_features, model.coef_.ravel())
        }
        stability_lookup = stability_table.set_index("feature")
        selected_records.append((outer_record.repeat, outer_record.fold, model_features))
        for feature in feature_names:
            in_ranking = feature in stability_lookup.index
            outer_feature_rows.append(
                {
                    "dataset": spec.key,
                    "scenario": scenario.name,
                    "outer_iter": outer_iter,
                    "outer_repeat": outer_record.repeat,
                    "outer_fold": outer_record.fold,
                    "feature": feature,
                    "eligible": feature in outer_eligible,
                    "rank": (
                        int(stability_lookup.loc[feature, "rank"]) if in_ranking else np.nan
                    ),
                    "eligible_freq_inner": (
                        float(stability_lookup.loc[feature, "eligible_freq"])
                        if in_ranking
                        else 0.0
                    ),
                    "nonzero_freq_inner": (
                        float(stability_lookup.loc[feature, "nonzero_freq"])
                        if in_ranking
                        else 0.0
                    ),
                    "mean_abs_coef_inner": (
                        float(stability_lookup.loc[feature, "mean_abs_coef"])
                        if in_ranking
                        else 0.0
                    ),
                    "selected_topk": feature in selected,
                    "model_used": feature in model_features,
                    "outer_coefficient": coefficient_map.get(feature, np.nan),
                    "k_selected": k_selected,
                    "n_model_features": len(model_features),
                }
            )

        train_counts = class_counts(y_train)
        test_counts = class_counts(y_test)
        parameter_rows.append(
            {
                "dataset": spec.key,
                "scenario": scenario.name,
                "outer_iter": outer_iter,
                "outer_repeat": outer_record.repeat,
                "outer_fold": outer_record.fold,
                "outer_splits_used": outer_plan.n_splits,
                "outer_repeats": outer_plan.n_repeats,
                "inner_splits_used": inner_plan.n_splits,
                "inner_repeats": inner_plan.n_repeats,
                "train_n": train_counts["n"],
                "train_positive": train_counts["positive"],
                "train_negative": train_counts["negative"],
                "test_n": test_counts["n"],
                "test_positive": test_counts["positive"],
                "test_negative": test_counts["negative"],
                "en_C": en_C,
                "en_l1_ratio": en_l1,
                "k_selected": k_selected,
                "l2_C": l2_C,
                "n_ranked_features": len(ranking),
                "n_model_features": len(model_features),
            }
        )
        print(
            f"[{label}] outer {outer_iter}/{len(outer_plan.records)} complete "
            f"(repeat={outer_record.repeat}, fold={outer_record.fold})"
        )

    predictions_long = pd.DataFrame(prediction_rows)
    expected_rows = n * OUTER_REPEATS
    if len(predictions_long) != expected_rows:
        raise RuntimeError(
            f"{label}: expected {expected_rows} repeat-level OOF rows, found {len(predictions_long)}"
        )
    per_id_repeat = predictions_long.groupby([ID_COL, "repeat"], sort=False).size()
    if not (per_id_repeat == 1).all() or len(per_id_repeat) != expected_rows:
        raise RuntimeError(f"{label}: each ID must be predicted once in every outer repeat")
    if not predictions_long["p"].between(0.0, 1.0).all():
        raise RuntimeError(f"{label}: invalid probability outside [0, 1]")

    oof = (
        predictions_long.groupby(ID_COL, sort=False)
        .agg(
            y_true=("y_true", "first"),
            p_oof=("p", "mean"),
            p_oof_sd=("p", "std"),
            oof_count=("p", "size"),
        )
        .reset_index()
    )
    expected_ids = ids.tolist()
    oof = oof.set_index(ID_COL).loc[expected_ids].reset_index()
    if not (oof["oof_count"] == OUTER_REPEATS).all():
        raise RuntimeError(f"{label}: incomplete averaged OOF predictions")
    if not np.array_equal(oof["y_true"].to_numpy(dtype=int), y):
        raise RuntimeError(f"{label}: OOF labels are misaligned with input IDs")

    repeat_metric_rows: List[Dict[str, Any]] = []
    for repeat, subset in predictions_long.groupby("repeat", sort=True):
        y_repeat = subset["y_true"].to_numpy(dtype=int)
        p_repeat = subset["p"].to_numpy(dtype=float)
        repeat_metric_rows.append(
            {
                "dataset": spec.key,
                "scenario": scenario.name,
                "repeat": int(repeat),
                **class_counts(y_repeat),
                "roc_auc": float(roc_auc_score(y_repeat, p_repeat)),
                "average_precision": float(average_precision_score(y_repeat, p_repeat)),
                "brier": float(brier_score_loss(y_repeat, p_repeat)),
            }
        )
    repeat_metrics = pd.DataFrame(repeat_metric_rows)

    y_oof = oof["y_true"].to_numpy(dtype=int)
    p_oof = oof["p_oof"].to_numpy(dtype=float)
    auc = float(roc_auc_score(y_oof, p_oof))
    ap = float(average_precision_score(y_oof, p_oof))
    brier = float(brier_score_loss(y_oof, p_oof))
    ci_low, ci_high = bootstrap_auc_ci(y_oof, p_oof, BOOTSTRAP_N, RANDOM_STATE)

    metrics = {
        "dataset": spec.key,
        "scenario": scenario.name,
        "analysis_role": spec.primary_role if scenario.name == "primary" else "sensitivity",
        "n_samples": n,
        "n_positive": counts["positive"],
        "n_negative": counts["negative"],
        "n_features_input": input_feature_count,
        "n_features": p,
        "n_features_globally_excluded": excluded_global_count,
        "positive_prevalence": float(np.mean(y_oof)),
        "roc_auc_oof": auc,
        "roc_auc_ci95_lo": ci_low,
        "roc_auc_ci95_hi": ci_high,
        "roc_auc_ci_method": "stratified patient bootstrap; descriptive, not CV-variance corrected",
        "average_precision_oof": ap,
        "pr_auc_oof": ap,
        "brier_oof": brier,
        "repeat_roc_auc_mean": float(repeat_metrics["roc_auc"].mean()),
        "repeat_roc_auc_sd": float(repeat_metrics["roc_auc"].std(ddof=1)),
        "repeat_average_precision_mean": float(repeat_metrics["average_precision"].mean()),
        "repeat_average_precision_sd": float(repeat_metrics["average_precision"].std(ddof=1)),
        "repeat_brier_mean": float(repeat_metrics["brier"].mean()),
        "repeat_brier_sd": float(repeat_metrics["brier"].std(ddof=1)),
        "outer_splits_used": outer_plan.n_splits,
        "outer_repeats": outer_plan.n_repeats,
    }

    outer_features = pd.DataFrame(outer_feature_rows)
    parameters = pd.DataFrame(parameter_rows)
    fold_counts = pd.DataFrame(fold_count_rows)
    k_curve = pd.concat(k_curve_frames, ignore_index=True)
    k_summary = (
        k_curve.groupby("k", sort=True)["mean_auc"]
        .agg(["mean", "std", "count"])
        .reset_index()
        .rename(
            columns={
                "mean": "mean_auc_across_outer",
                "std": "std_auc_across_outer",
                "count": "n_outer_fits",
            }
        )
    )
    k_summary["se_auc_across_outer"] = (
        k_summary["std_auc_across_outer"] / np.sqrt(k_summary["n_outer_fits"])
    )

    feature_frequency = (
        outer_features.groupby("feature", sort=False)
        .agg(
            eligible_count=("eligible", "sum"),
            selected_topk_count=("selected_topk", "sum"),
            model_used_count=("model_used", "sum"),
            coefficient_mean=("outer_coefficient", "mean"),
            coefficient_median=("outer_coefficient", "median"),
        )
        .reset_index()
    )
    denominator = float(len(outer_plan.records))
    feature_frequency["eligible_rate"] = feature_frequency["eligible_count"] / denominator
    feature_frequency["selected_topk_rate"] = feature_frequency["selected_topk_count"] / denominator
    feature_frequency["selected_rate"] = feature_frequency["model_used_count"] / denominator
    positive_counts = (
        outer_features.assign(coef_positive=outer_features["outer_coefficient"] > 0)
        .groupby("feature", sort=False)["coef_positive"]
        .sum()
        .rename("coefficient_positive_count")
        .reset_index()
    )
    feature_frequency = feature_frequency.merge(positive_counts, on="feature", how="left")
    feature_frequency["coefficient_positive_rate_among_used"] = np.where(
        feature_frequency["model_used_count"] > 0,
        feature_frequency["coefficient_positive_count"] / feature_frequency["model_used_count"],
        np.nan,
    )
    feature_frequency = feature_frequency.sort_values(
        ["selected_rate", "feature"], ascending=[False, True]
    ).reset_index(drop=True)

    selected_sets = [features for _, _, features in selected_records]
    stability = nogueira_stability(selected_sets, feature_names)
    jaccard = pairwise_jaccard_table(selected_records)
    stability.update(summarize_jaccard(jaccard))
    stability.update({"dataset": spec.key, "scenario": scenario.name})
    stability_summary = pd.DataFrame([stability])

    # Final full-data signature. This is not an independent validation result.
    final_inner = make_cv_plan(
        y,
        target_splits=INNER_SPLITS_TARGET,
        n_repeats=INNER_REPEATS,
        context=f"{label} final-inner",
        log_lines=log_lines,
        random_state=RANDOM_STATE,
    )
    final_en_C, final_en_l1, _ = tune_elasticnet(
        X,
        y,
        feature_names,
        final_inner,
        dataset_name=label,
        outer_iter=0,
        log_lines=log_lines,
        feature_rules=rules,
        imputation=scenario.imputation,
        warning_rows=warning_rows,
    )
    final_ranking, _ = stability_select_and_rank(
        X,
        y,
        feature_names,
        final_inner,
        en_C=final_en_C,
        en_l1=final_en_l1,
        dataset_name=label,
        outer_iter=0,
        log_lines=log_lines,
        feature_rules=rules,
        imputation=scenario.imputation,
        warning_rows=warning_rows,
    )
    full_eligibility = _make_preprocessor(feature_names, rules, scenario.imputation).fit(X)
    full_eligible = set(full_eligibility.kept_features_)
    final_ranking = [feature for feature in final_ranking if feature in full_eligible]
    final_k, final_l2_C, _ = select_k_by_1se(
        X,
        y,
        final_ranking,
        final_inner,
        dataset_name=label,
        outer_iter=0,
        log_lines=log_lines,
        feature_rules=rules,
        imputation=scenario.imputation,
        max_k=scenario.max_k,
        warning_rows=warning_rows,
    )
    final_selected = final_ranking[:final_k]
    final_preprocessor = _make_preprocessor(
        final_selected, rules, scenario.imputation
    ).fit(X[final_selected])
    final_matrix = final_preprocessor.transform(X[final_selected])
    final_model = fit_logreg_with_refit(
        final_matrix,
        y,
        kind="l2",
        C=final_l2_C,
        l1_ratio=None,
        dataset_name=label,
        stage="L2-finalfit",
        log_lines=log_lines,
        warning_rows=warning_rows,
    )
    final_model_features = final_preprocessor.kept_features_
    final_feature_list = pd.DataFrame(
        {
            "rank": np.arange(1, len(final_selected) + 1, dtype=int),
            "feature": final_selected,
            "model_used": [feature in final_model_features for feature in final_selected],
        }
    )
    final_coefficients = pd.DataFrame(
        {
            "feature": final_model_features,
            "coefficient": final_model.coef_.ravel(),
        }
    )
    final_coefficients = final_coefficients.sort_values(
        "coefficient", key=lambda values: np.abs(values), ascending=False
    ).reset_index(drop=True)
    final_coefficients.loc[len(final_coefficients)] = {
        "feature": "INTERCEPT",
        "coefficient": float(final_model.intercept_.ravel()[0]),
    }
    final_preprocessing = final_preprocessor.audit_table()
    final_spec = {
        "dataset": spec.key,
        "scenario": scenario.name,
        "statement": "Final full-data fitted signature; not an independent validation result",
        "selected_features_topk": final_selected,
        "model_features_after_preprocessing": final_model_features,
        "final_en_C": final_en_C,
        "final_en_l1_ratio": final_en_l1,
        "final_k": final_k,
        "final_l2_C": final_l2_C,
        "class_weight": "balanced",
        "deployment_threshold": None,
    }

    # Plot-source coordinates only; final scientific plotting is intentionally
    # kept separate from model training.
    fpr, tpr, roc_threshold = roc_curve(y_oof, p_oof)
    roc_coordinates = pd.DataFrame(
        {"fpr": fpr, "tpr": tpr, "threshold": roc_threshold}
    )
    precision, recall, pr_threshold = precision_recall_curve(y_oof, p_oof)
    pr_coordinates = pd.DataFrame(
        {
            "precision": precision,
            "recall": recall,
            "threshold": np.append(pr_threshold, np.nan),
        }
    )

    # Persist only after all modeling and integrity checks have completed.
    to_csv_utf8sig(predictions_long, output_dir / "outer_predictions_long.csv")
    to_csv_utf8sig(oof, output_dir / "oof_predictions.csv")
    to_csv_utf8sig(repeat_metrics, output_dir / "repeat_metrics.csv")
    to_csv_utf8sig(pd.DataFrame([metrics]), output_dir / "metrics_summary.csv")
    to_csv_utf8sig(parameters, output_dir / "outer_parameters.csv")
    to_csv_utf8sig(parameters, output_dir / "tuning_summary.csv")
    to_csv_utf8sig(fold_counts, output_dir / "fold_counts.csv")
    to_csv_utf8sig(outer_features, output_dir / "outer_features_long.csv")
    to_csv_utf8sig(feature_frequency, output_dir / "feature_selection_frequency.csv")
    to_csv_utf8sig(stability_summary, output_dir / "feature_stability_summary.csv")
    to_csv_utf8sig(jaccard, output_dir / "pairwise_jaccard.csv")
    to_csv_utf8sig(k_curve, output_dir / "k_curve_per_outer.csv")
    to_csv_utf8sig(k_summary, output_dir / "k_curve_summary.csv")
    to_csv_utf8sig(bundle.feature_manifest, output_dir / "feature_manifest.csv")
    to_csv_utf8sig(roc_coordinates, output_dir / "roc_curve.csv")
    to_csv_utf8sig(pr_coordinates, output_dir / "pr_curve.csv")
    to_tsv_utf8sig(final_feature_list, output_dir / "final_feature_list.tsv")
    to_tsv_utf8sig(final_coefficients, output_dir / "final_coefficients.tsv")
    to_csv_utf8sig(final_preprocessing, output_dir / "final_preprocessing.csv")
    warning_table = pd.DataFrame(
        warning_rows,
        columns=["dataset", "stage", "warning_type", "message"],
    )
    to_csv_utf8sig(warning_table, output_dir / "warnings.csv")
    write_json(output_dir / "final_model_spec.json", final_spec)
    write_json(
        output_dir / "settings_used.json",
        _scenario_settings_payload(spec, scenario, outer_plan.n_splits),
    )

    elapsed = time.time() - started
    script_path = Path(__file__).resolve()
    receipt = {
        "status": "completed",
        "started_utc": started_iso,
        "completed_utc": datetime.now(timezone.utc).isoformat(),
        "elapsed_seconds": elapsed,
        "dataset": spec.key,
        "scenario": scenario.name,
        "invocation_scenario": ACTIVE_SCENARIO,
        "invocation_scenario_metadata": ACTIVE_SCENARIO_METADATA,
        "analysis_role": metrics["analysis_role"],
        "input_path": str(spec.path),
        "input_sha256": bundle.input_sha256,
        "input_sheet": scenario.sheet_name,
        "id_column": ID_COL,
        "label_column": bundle.label_col,
        "label_mapping": {"Responder": 1, "Non-responder": 0, "numeric": "0/1"},
        "script_path": str(script_path),
        "script_sha256": sha256_file(script_path),
        **_selector_provenance(),
        "python_version": platform.python_version(),
        "package_versions": {
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "scikit_learn": sklearn.__version__,
            "openpyxl": openpyxl.__version__,
        },
        "sample_counts": counts,
        "input_feature_count": input_feature_count,
        "feature_count": p,
        "globally_excluded_feature_count": excluded_global_count,
        "outer_splits_used": outer_plan.n_splits,
        "outer_repeats": outer_plan.n_repeats,
        "warning_count": len(warning_rows),
        "metrics": metrics,
        "feature_stability": stability,
        "final_model": final_spec,
        "output_dir": str(output_dir),
    }
    write_json(output_dir / "run_receipt.json", receipt)
    log_lines.append(f"[DONE] elapsed_seconds={elapsed:.1f}, output={output_dir}")
    write_text(output_dir / "run_log.txt", "\n".join(log_lines) + "\n")

    print(f"[DONE] {label} -> {output_dir}")
    return RunResult(
        dataset_key=spec.key,
        scenario=scenario.name,
        output_dir=output_dir,
        outer_splits_used=outer_plan.n_splits,
        max_selected_k=int(parameters["k_selected"].max()),
        primary_k_ceiling_hit=bool((parameters["k_selected"] >= scenario.max_k).any()),
        metrics=metrics,
    )


# =============================================================================
# SCENARIO SCHEDULING
# =============================================================================

def _intersect_with_global_targets(keys: Sequence[str]) -> List[str]:
    global_selected = set(_profile_targets("datasets_to_run", DATASETS_TO_RUN, keys))
    return [key for key in keys if key in global_selected]


def _primary_scenario(spec: DatasetSpec) -> ScenarioSpec:
    profile = _SCENARIO_PROFILES[ACTIVE_SCENARIO]
    return ScenarioSpec(
        name=str(profile["primary_scenario_name"]),
        sheet_name=spec.primary_sheet,
        imputation="half_minimum",
        max_k=PRIMARY_MAX_K,
        outer_splits_target=OUTER_SPLITS_TARGET,
        source=str(profile["primary_scenario_source"]),
    )


def _sensitivity_scenarios(
    specs: Mapping[str, DatasetSpec],
    primary_results: Mapping[str, RunResult],
) -> List[Tuple[DatasetSpec, ScenarioSpec]]:
    available = _intersect_with_global_targets(list(specs))
    scheduled: List[Tuple[DatasetSpec, ScenarioSpec]] = []

    if RUN_SD_SENSITIVITY:
        for key in _profile_targets("sd_targets", SD_TARGETS, available):
            spec = specs[key]
            if spec.dataset_id != "dataset1":
                raise ValueError(f"SD sensitivity sheets are supported only for Dataset 1: {key}")
            scheduled.extend(
                [
                    (
                        spec,
                        ScenarioSpec(
                            name="sd_as_nr",
                            sheet_name=SD_AS_NR_SHEET,
                            imputation="half_minimum",
                            max_k=PRIMARY_MAX_K,
                            outer_splits_target=OUTER_SPLITS_TARGET,
                            source="sd_sensitivity",
                        ),
                    ),
                    (
                        spec,
                        ScenarioSpec(
                            name="sd_as_r",
                            sheet_name=SD_AS_R_SHEET,
                            imputation="half_minimum",
                            max_k=PRIMARY_MAX_K,
                            outer_splits_target=OUTER_SPLITS_TARGET,
                            source="sd_sensitivity",
                        ),
                    ),
                ]
            )

    if RUN_KNN_SENSITIVITY:
        for key in _profile_targets("knn_targets", KNN_TARGETS, available):
            spec = specs[key]
            scheduled.append(
                (
                    spec,
                    ScenarioSpec(
                        name="knn",
                        sheet_name=spec.primary_sheet,
                        imputation="knn",
                        max_k=PRIMARY_MAX_K,
                        outer_splits_target=OUTER_SPLITS_TARGET,
                        source="imputation_sensitivity",
                    ),
                )
            )

    if RUN_OUTER3_SENSITIVITY:
        outer3_available = list(primary_results) if OUTER3_TARGETS == "ALL_PRIMARY" else available
        for key in _profile_targets("outer3_targets", OUTER3_TARGETS, outer3_available):
            spec = specs[key]
            scheduled.append(
                (
                    spec,
                    ScenarioSpec(
                        name="outer3",
                        sheet_name=spec.primary_sheet,
                        imputation="half_minimum",
                        max_k=PRIMARY_MAX_K,
                        outer_splits_target=OUTER3_SPLITS_TARGET,
                        source="fold_sensitivity",
                    ),
                )
            )

    if RUN_KMAX_SENSITIVITY:
        target_set = set(_profile_targets("kmax_targets", KMAX_EXPLICIT_TARGETS, available))
        if AUTO_KMAX_FROM_PRIMARY:
            target_set.update(
                key
                for key, result in primary_results.items()
                if result.primary_k_ceiling_hit
            )
        for key in available:
            if key not in target_set:
                continue
            spec = specs[key]
            scheduled.append(
                (
                    spec,
                    ScenarioSpec(
                        name="kmax30",
                        sheet_name=spec.primary_sheet,
                        imputation="half_minimum",
                        max_k=KMAX_SENSITIVITY_MAX_K,
                        outer_splits_target=OUTER_SPLITS_TARGET,
                        source="kmax_sensitivity",
                    ),
                )
            )

    unique: Dict[Tuple[str, str], Tuple[DatasetSpec, ScenarioSpec]] = {}
    for spec, scenario in scheduled:
        unique[(spec.key, scenario.name)] = (spec, scenario)
    return list(unique.values())


def _master_manifest_row(result: RunResult) -> Dict[str, Any]:
    return {
        "dataset": result.dataset_key,
        "scenario": result.scenario,
        "invocation_scenario": ACTIVE_SCENARIO,
        "outer_splits_used": result.outer_splits_used,
        "max_selected_k": result.max_selected_k,
        "primary_k_ceiling_hit": result.primary_k_ceiling_hit,
        "roc_auc_oof": result.metrics["roc_auc_oof"],
        "average_precision_oof": result.metrics["average_precision_oof"],
        "brier_oof": result.metrics["brier_oof"],
        "output_dir": str(result.output_dir),
    }


def main() -> None:
    specs = discover_datasets()
    available = _intersect_with_global_targets(list(specs))
    root = OUTPUT_ROOT / RUN_TAG if RUN_TAG else OUTPUT_ROOT
    root.mkdir(parents=True, exist_ok=True)
    coverage_path = root / "input_coverage.csv"
    if coverage_path.exists() and not ALLOW_OVERWRITE:
        raise FileExistsError(f"Existing invocation coverage will not be overwritten: {coverage_path}")
    coverage = pd.DataFrame(coverage_rows(specs, available))
    coverage["invocation_scenario"] = ACTIVE_SCENARIO
    coverage["invocation_status"] = "scheduled" if available else "no_applicable_inputs"
    coverage["input_file"] = coverage["dataset"].map(
        {key: str(spec.path) for key, spec in specs.items()}
    ).fillna("")
    to_csv_utf8sig(coverage, coverage_path)
    if not available:
        print(f"No supplied inputs apply to scenario {ACTIVE_SCENARIO!r}; "
              f"coverage recorded in {root / 'input_coverage.csv'}")
        return

    results: List[RunResult] = []
    primary_results: Dict[str, RunResult] = {}

    if RUN_PRIMARY:
        for key in _profile_targets("primary_targets", PRIMARY_TARGETS, available):
            result = run_nested_cv(specs[key], _primary_scenario(specs[key]))
            results.append(result)
            primary_results[key] = result

    for spec, scenario in _sensitivity_scenarios(specs, primary_results):
        results.append(run_nested_cv(spec, scenario))

    completed_keys = {result.dataset_key for result in results}
    coverage["selection_status"] = coverage["dataset"].map(
        lambda key: "selected" if key in completed_keys else "not_selected"
    )
    if not results:
        coverage["invocation_status"] = "no_analyses_scheduled"
        to_csv_utf8sig(coverage, coverage_path)
        print(f"No analyses were scheduled for scenario {ACTIVE_SCENARIO!r}; "
              f"check target selectors and enabled analyses. Coverage: {coverage_path}")
        return

    manifest = pd.DataFrame([_master_manifest_row(result) for result in results])
    if not manifest.empty:
        to_csv_utf8sig(manifest, root / "run_manifest.csv")
    coverage["invocation_status"] = "completed"
    to_csv_utf8sig(coverage, coverage_path)
    print("=" * 88)
    print(f"All requested analyses completed: {len(results)} run(s)")
    print(f"Manifest: {root / 'run_manifest.csv'}")


def _parse_cli_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Run the leakage-aware ICI prediction pipeline using one recorded "
            "primary or sensitivity-analysis configuration."
        )
    )
    parser.add_argument(
        "--scenario",
        choices=PUBLIC_SCENARIOS,
        default="primary",
        help="Analysis configuration to run (default: primary).",
    )
    parser.add_argument(
        "--datasets", nargs="+", default=None, metavar="DATASET_KEY",
        help="Explicit canonical inputs; unknown or unavailable keys are errors.",
    )
    return parser.parse_args(argv)


def configure_cli(args: argparse.Namespace) -> None:
    global DATASETS_TO_RUN, EXPLICIT_DATASET_SELECTION
    configure_scenario(args.scenario)
    if args.datasets is not None:
        # Restrict requests to the scenario's intended study matrices.
        profile_selector = _SCENARIO_PROFILES[ACTIVE_SCENARIO]["datasets_to_run"]
        eligible = select_available(profile_selector, KNOWN_DATASET_KEYS)
        select_available(args.datasets, eligible)
        DATASETS_TO_RUN = list(args.datasets)
        EXPLICIT_DATASET_SELECTION = True


if __name__ == "__main__":
    configure_cli(_parse_cli_args())
    main()
