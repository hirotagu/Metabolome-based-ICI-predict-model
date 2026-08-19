#!/usr/bin/env python3
"""Create the 17 canonical analysis workbooks from Supplementary Data 1.

The source workbook is the analysis-ready public dataset.  Each data sheet is
validated, standardized to ``id``/``category`` followed by numeric features,
and written as a one-sheet workbook named ``datasetN_pre/post1/post2.xlsx``.
Existing output content is never overwritten.
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import tempfile
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# =============================================================================
# USER SETTINGS
# =============================================================================

SCRIPT_DIR = Path(__file__).resolve().parent
INPUT_XLSX = SCRIPT_DIR / "Supplementary Data 1.xlsx"
OUTPUT_DIR = SCRIPT_DIR / "raw"

ID_COLUMN = "id"
LABEL_COLUMN = "category"
SUSPECT_DATASET5_POST2_ID = "CA209025-43-743"

# sheet, canonical key, N, responders/positive, non-responders/negative, features
SHEET_PLAN: Tuple[Tuple[str, str, int, int, int, int], ...] = (
    ("Sheet1", "dataset1_pre", 51, 35, 16, 128),
    ("Sheet2", "dataset2_pre", 47, 25, 22, 124),
    ("Sheet3", "dataset3_pre", 77, 57, 20, 202),
    ("Sheet4", "dataset4_pre", 72, 49, 23, 107),
    ("Sheet5", "dataset5_pre", 381, 199, 182, 202),
    ("Sheet6", "dataset6_pre", 17, 13, 4, 168),
    ("Sheet7", "dataset7_pre", 39, 21, 18, 105),
    ("Sheet8", "dataset1_post1", 51, 35, 16, 128),
    ("Sheet9", "dataset2_post1", 46, 24, 22, 124),
    ("Sheet10", "dataset3_post1", 71, 55, 16, 202),
    ("Sheet11", "dataset4_post1", 64, 46, 18, 107),
    ("Sheet12", "dataset5_post1", 95, 60, 35, 202),
    ("Sheet13", "dataset6_post1", 18, 13, 5, 168),
    ("Sheet14", "dataset2_post2", 47, 25, 22, 124),
    ("Sheet15", "dataset3_post2", 61, 51, 10, 202),
    ("Sheet16", "dataset4_post2", 61, 47, 14, 107),
    ("Sheet17", "dataset5_post2", 314, 183, 131, 202),
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normalized_header(value: object) -> str:
    return " ".join(str(value).strip().split()).casefold()


def _normalize_id(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)):
        number = float(value)
        if np.isfinite(number) and number.is_integer():
            return str(int(number))
    return str(value).strip()


def _binary_label(value: object) -> int:
    if pd.isna(value):
        raise ValueError("missing label")
    if isinstance(value, (bool, np.bool_)):
        return int(value)
    if isinstance(value, (int, np.integer)) and int(value) in {0, 1}:
        return int(value)
    if (
        isinstance(value, (float, np.floating))
        and np.isfinite(value)
        and float(value) in {0.0, 1.0}
    ):
        return int(value)
    token = _normalized_header(value).replace("_", "-")
    mapping = {
        "responder": 1,
        "response": 1,
        "r": 1,
        "positive": 1,
        "non-responder": 0,
        "nonresponder": 0,
        "nr": 0,
        "negative": 0,
    }
    if token not in mapping:
        raise ValueError(f"unknown label {value!r}")
    return mapping[token]


def _raw_headers(workbook_path: Path, sheet_name: str) -> List[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        raw = list(
            next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                (),
            )
        )
    finally:
        workbook.close()
    while raw and raw[-1] is None:
        raw.pop()
    headers = ["" if value is None else str(value).strip() for value in raw]
    if any(not value for value in headers):
        positions = [index + 1 for index, value in enumerate(headers) if not value]
        raise ValueError(f"{sheet_name}: blank header(s) in columns {positions}")
    normalized = [_normalized_header(value) for value in headers]
    duplicates = sorted({value for value in normalized if normalized.count(value) > 1})
    if duplicates:
        raise ValueError(f"{sheet_name}: duplicate header(s): {duplicates}")
    return headers


def _read_validate_sheet(
    workbook_path: Path,
    sheet_name: str,
    dataset_key: str,
    expected_n: int,
    expected_positive: int,
    expected_negative: int,
    expected_features: int,
) -> Tuple[pd.DataFrame, Dict[str, object]]:
    headers = _raw_headers(workbook_path, sheet_name)
    frame = pd.read_excel(
        workbook_path,
        sheet_name=sheet_name,
        engine="openpyxl",
        na_values=["NA", "N/A", "na", "n/a"],
        keep_default_na=True,
    ).dropna(how="all").reset_index(drop=True)
    frame.columns = [str(value).strip() for value in frame.columns]
    if frame.columns.tolist() != headers:
        raise ValueError(f"{sheet_name}: header mismatch after reading the workbook")
    if frame.shape[1] < 3:
        raise ValueError(f"{sheet_name}: expected id, label, and at least one feature")
    if frame.columns[0] != ID_COLUMN:
        raise ValueError(
            f"{sheet_name}: first column must be {ID_COLUMN!r}, found {frame.columns[0]!r}"
        )
    if frame.columns[1] not in {"category", "Responder"}:
        raise ValueError(
            f"{sheet_name}: second column must be 'category' or 'Responder', "
            f"found {frame.columns[1]!r}"
        )

    ids = frame.iloc[:, 0].map(_normalize_id)
    if (ids == "").any():
        rows = (np.flatnonzero((ids == "").to_numpy()) + 2).tolist()[:10]
        raise ValueError(f"{sheet_name}: blank ID(s) in Excel rows {rows}")
    if ids.duplicated().any():
        duplicates = ids.loc[ids.duplicated(keep=False)].unique().tolist()[:10]
        raise ValueError(f"{sheet_name}: duplicate ID(s): {duplicates}")

    labels: List[int] = []
    label_errors: List[str] = []
    for row_number, value in enumerate(frame.iloc[:, 1].tolist(), start=2):
        try:
            labels.append(_binary_label(value))
        except ValueError as exc:
            label_errors.append(f"row {row_number}: {exc}")
    if label_errors:
        raise ValueError(f"{sheet_name}: invalid labels: {'; '.join(label_errors[:10])}")

    feature_names = frame.columns[2:].tolist()
    reserved = {_normalized_header(ID_COLUMN), "category", "responder"}
    reserved_features = [
        feature for feature in feature_names if _normalized_header(feature) in reserved
    ]
    if reserved_features:
        raise ValueError(
            f"{sheet_name}: reserved ID/label name(s) in feature block: "
            f"{reserved_features}"
        )
    numeric_features: Dict[str, pd.Series] = {}
    conversion_errors: List[str] = []
    for feature in feature_names:
        original = frame[feature]
        numeric = pd.to_numeric(original, errors="coerce")
        bad = original.notna() & numeric.isna()
        if bad.any():
            examples = original.loc[bad].astype(str).unique().tolist()[:5]
            conversion_errors.append(f"{feature}: {examples}")
        numeric_features[feature] = numeric.astype(float)
    if conversion_errors:
        raise ValueError(
            f"{sheet_name}: non-numeric feature values: "
            + "; ".join(conversion_errors[:10])
        )

    output = pd.DataFrame(
        {
            ID_COLUMN: ids,
            LABEL_COLUMN: labels,
            **{feature: numeric_features[feature] for feature in feature_names},
        }
    )

    observed = {
        "n": int(len(output)),
        "positive": int(output[LABEL_COLUMN].eq(1).sum()),
        "negative": int(output[LABEL_COLUMN].eq(0).sum()),
        "features": int(len(feature_names)),
    }
    expected = {
        "n": int(expected_n),
        "positive": int(expected_positive),
        "negative": int(expected_negative),
        "features": int(expected_features),
    }
    if observed != expected:
        raise ValueError(
            f"{sheet_name}/{dataset_key}: QC counts differ; "
            f"expected={expected}, observed={observed}"
        )
    suspect_absent = SUSPECT_DATASET5_POST2_ID not in set(ids)
    if dataset_key == "dataset5_post2" and not suspect_absent:
        raise ValueError(
            f"{sheet_name}: excluded suspect ID is present: "
            f"{SUSPECT_DATASET5_POST2_ID}"
        )

    qc: Dict[str, object] = {
        "source_sheet": sheet_name,
        "dataset": dataset_key,
        "output_file": f"{dataset_key}.xlsx",
        "n": observed["n"],
        "positive": observed["positive"],
        "negative": observed["negative"],
        "features": observed["features"],
        "id_unique": True,
        "binary_label": True,
        "feature_values_numeric_or_missing": True,
        "suspect_dataset5_post2_id_absent": suspect_absent,
    }
    return output, qc


def _validate_written_workbook(
    path: Path,
    expected_rows: int,
    expected_columns: int,
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise RuntimeError(
                f"{path.name}: expected exactly one sheet named Sheet1, "
                f"found {workbook.sheetnames}"
            )
        worksheet = workbook["Sheet1"]
        if worksheet.max_row != expected_rows + 1:
            raise RuntimeError(
                f"{path.name}: expected {expected_rows} data rows, "
                f"found {worksheet.max_row - 1}"
            )
        if worksheet.max_column != expected_columns:
            raise RuntimeError(
                f"{path.name}: expected {expected_columns} columns, "
                f"found {worksheet.max_column}"
            )
        first = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if tuple(first[:2]) != (ID_COLUMN, LABEL_COLUMN):
            raise RuntimeError(f"{path.name}: invalid output header {first[:2]}")
    finally:
        workbook.close()


def prepare_inputs(input_xlsx: Path, output_dir: Path) -> pd.DataFrame:
    input_xlsx = input_xlsx.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    if not input_xlsx.is_file():
        raise FileNotFoundError(f"Input workbook not found: {input_xlsx}")

    workbook = load_workbook(input_xlsx, read_only=True, data_only=True)
    try:
        expected_sheets = ["README", *(item[0] for item in SHEET_PLAN)]
        if workbook.sheetnames != expected_sheets:
            raise ValueError(
                "Supplementary Data 1 sheet structure differs from the published "
                f"mapping; expected={expected_sheets}, found={workbook.sheetnames}"
            )
    finally:
        workbook.close()

    planned_paths = [output_dir / f"{item[1]}.xlsx" for item in SHEET_PLAN]
    existing_targets = [path for path in planned_paths if path.exists()]
    if existing_targets:
        raise FileExistsError(
            "Existing analysis inputs will not be overwritten: "
            + ", ".join(path.name for path in existing_targets)
        )
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(
            f"Output directory already contains files and will not be modified: {output_dir}"
        )

    validated: List[Tuple[str, pd.DataFrame]] = []
    qc_rows: List[Dict[str, object]] = []
    for sheet, key, n, positive, negative, features in SHEET_PLAN:
        frame, qc = _read_validate_sheet(
            input_xlsx,
            sheet,
            key,
            n,
            positive,
            negative,
            features,
        )
        validated.append((key, frame))
        qc_rows.append(qc)

    output_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix="prepare_inputs_",
        dir=output_dir.parent,
    ) as temporary:
        stage = Path(temporary) / output_dir.name
        stage.mkdir()
        for key, frame in validated:
            path = stage / f"{key}.xlsx"
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                frame.to_excel(writer, sheet_name="Sheet1", index=False)
            _validate_written_workbook(
                path,
                expected_rows=len(frame),
                expected_columns=frame.shape[1],
            )
        qc_table = pd.DataFrame(qc_rows)
        qc_table.insert(0, "source_workbook", input_xlsx.name)
        qc_table.insert(1, "source_sha256", _sha256(input_xlsx))
        qc_table["output_sha256"] = [
            _sha256(stage / f"{key}.xlsx") for key, _ in validated
        ]
        qc_table.to_csv(stage / "input_qc.csv", index=False, encoding="utf-8-sig")

        if output_dir.exists():
            output_dir.rmdir()  # Preflight guarantees that it is empty.
        shutil.move(str(stage), str(output_dir))

    print(qc_table[["dataset", "n", "positive", "negative", "features"]].to_string(index=False))
    print(f"Prepared {len(qc_table)} canonical workbooks: {output_dir}")
    return qc_table


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Prepare the 17 canonical analysis inputs from Supplementary Data 1."
    )
    parser.add_argument("--input", type=Path, default=INPUT_XLSX)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    prepare_inputs(args.input, args.output_dir)


if __name__ == "__main__":
    main()
